"""2026-10-17 (R52): two follow-ups to an analysis that used to re-run the whole analysis.

"Show the 135 high-risk customers as a list"  -> the tier's rows, highest score first, with the signals each row
trips, and an Excel/CSV download. Nothing is recomputed: the rows were scored out of fold when the analysis ran.

"What if churn fell by a quarter in the low-usage band?" -> plain arithmetic on the analysis's own groups. The
assumptions are stated as assumptions (the cut applies only to that group, everyone else keeps their past rate,
rates are from the past period), never presented as findings, and every number comes from the analysis.
"""
from __future__ import annotations

import csv
import io
import re

from . import analysis as an_mod

LIST_Q = re.compile(
    r"\b(list|export|download|show|give|send|get|pull|see)\b.{0,60}\b(high[- ]?risk|at[- ]?risk|riskiest|medium[- ]?risk|most likely to (leave|churn|cancel))\b"
    r"|\b(high[- ]?risk|at[- ]?risk|medium[- ]?risk)\b.{0,50}\b(as|in|into) (a |an )?(list|table|spreadsheet|excel|csv|sheet)\b"
    r"|\b(the )?\d[\d,]* (high[- ]?risk|at[- ]?risk)\b", re.I)
WHATIF_Q = re.compile(r"\bwhat if\b|\bwhat would happen if\b|\bif we (could )?(cut|reduce|lower|halve|drop|bring down)\b|\bsimulate\b|\bscenario\b", re.I)

_FRACTIONS = [(r"\bhalf\b|\bhalve[ds]?\b|\bin half\b", 0.5), (r"\ba quarter\b|\bone quarter\b|\bquarter\b", 0.25), (r"\ba third\b|\bone third\b", 1 / 3),
              (r"\btwo thirds\b", 2 / 3), (r"\ba fifth\b", 0.2), (r"\ba tenth\b", 0.1)]
_USAGE = ("trans", "usage", "use", "spend", "amt", "amount", "activ", "util", "purchase", "order", "visit", "login", "session", "engag")
_SYN = {"usage": _USAGE, "use": _USAGE, "activity": _USAGE + ("inactiv",), "active": _USAGE, "engagement": _USAGE, "spend": ("amt", "amount", "spend", "revenue"),
        "inactive": ("inactiv",), "inactivity": ("inactiv",), "contact": ("contact", "call"), "tenure": ("month", "tenure", "book"),
        "products": ("relationship", "product"), "balance": ("bal", "revolv"), "income": ("income",), "age": ("age",), "card": ("card",)}


def wants_list(text: str) -> bool:
    return bool(LIST_Q.search(text or ""))


def wants_whatif(text: str) -> bool:
    return bool(WHATIF_Q.search(text or ""))


# ---- the list ---------------------------------------------------------------------------------

def tier_of(text: str) -> str:
    t = (text or "").lower()
    if re.search(r"\bmedium[- ]?risk\b", t):
        return "medium"
    if re.search(r"\b(all|every)\b.{0,20}\b(at[- ]?risk)\b|\bhigh and medium\b|\bmedium and high\b", t):
        return "both"
    return "high"


def build_list(a: dict, rows_by_tier: dict | None, text: str) -> dict:
    """{title, tier, rule, columns, rows, total, asked, note} from the analysis and its stored rows."""
    tiers = {t["key"]: t for t in (a.get("tiers") or [])}
    which = tier_of(text)
    keys = ["high", "medium"] if which == "both" else [which]
    rows: list[dict] = []
    for k in keys:
        for r in (rows_by_tier or {}).get(k) or []:
            rows.append(r | {"tier": tiers.get(k, {}).get("label", k)})
    total = sum(int(tiers.get(k, {}).get("count") or 0) for k in keys)
    ent = (a.get("table") or {}).get("entities") or "rows"
    label = " and ".join(tiers.get(k, {}).get("label", k).lower().replace(" risk", "-risk") for k in keys)
    cols = [c for c in ((a.get("risk_list") or {}).get("columns") or ["id", "score"])]
    id_label = (a.get("risk_list") or {}).get("id_column") or a.get("id_column") or "Row"  # the column's own name, as in the file
    asked = None
    m = re.search(r"\b(\d[\d,]*)\b", text or "")
    if m:
        try:
            asked = int(m.group(1).replace(",", ""))
        except ValueError:
            asked = None
    note = None
    if asked and asked != total and total:
        note = f"You asked for {asked:,}; the {label} tier holds {total:,} {ent} in this analysis."
    rule = " · ".join(tiers.get(k, {}).get("rule", "") for k in keys)
    return {"title": f"{label.capitalize()} {ent}", "tier": which, "rule": rule, "id_label": id_label,
            "columns": ["id", "score", "tier"] + [c for c in cols if c not in ("id", "score")], "rows": rows, "total": total,
            "shown": len(rows), "asked": asked, "note": note, "entities": ent}


def list_text(lst: dict) -> str:
    n, ent = lst["total"], lst["entities"]
    if not n:
        return f"No current {ent} are in that tier, so there is nothing to list."
    head = f"Here are the {n:,} {lst['title'].lower()}, highest score first"
    if lst["shown"] < n:
        head += f" (the first {lst['shown']:,})"
    s = head + ". Each row shows the signals that put it there; download it as Excel to work through it."
    return (lst["note"] + " " + s) if lst.get("note") else s


def _cell(v):
    if isinstance(v, list):
        return "; ".join(map(str, v))
    return v


def _head(lst: dict, c: str) -> str:
    if c == "id":
        return str(lst.get("id_label") or "Row")
    return c[:1].upper() + c[1:]


def list_csv(lst: dict) -> bytes:
    buf = io.StringIO()
    cols = lst["columns"] + ["signals"]
    w = csv.writer(buf)
    w.writerow([_head(lst, c) for c in cols])
    for r in lst["rows"]:
        w.writerow([_cell(r.get(c, "")) for c in cols])
    return buf.getvalue().encode("utf-8-sig")


def list_xlsx(lst: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = (lst["title"] or "List")[:31]
    cols = lst["columns"] + ["signals"]
    heads = [_head(lst, c) for c in cols]
    ws.append(heads)
    for i, _ in enumerate(heads, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B0F0E")
        cell.alignment = Alignment(vertical="center")
    for r in lst["rows"]:
        ws.append([_cell(r.get(c, "")) for c in cols])
    for i, c in enumerate(cols, 1):
        width = max(len(str(heads[i - 1])), *(len(str(_cell(r.get(c, "")))) for r in lst["rows"][:200])) if lst["rows"] else len(str(heads[i - 1]))
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    sc = cols.index("score") + 1
    for row in ws.iter_rows(min_row=2, min_col=sc, max_col=sc):
        for cell in row:
            cell.number_format = "0.00"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    meta = wb.create_sheet("About")
    meta.append(["What", lst["title"]])
    meta.append(["Rule", lst["rule"]])
    meta.append(["Rows", lst["total"]])
    meta.append(["Score", "Out-of-fold model score: each row was scored by a model that never saw it."])
    meta.column_dimensions["A"].width = 12
    meta.column_dimensions["B"].width = 80
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---- the what-if --------------------------------------------------------------------------------

def _change(text: str) -> tuple[str, float] | None:
    """("relative", 0.25) for "by a quarter / by 25%", ("points", 5) for "by 5 points"."""
    t = (text or "").lower()
    m = re.search(r"(\d+(?:\.\d+)?)\s*(?:percentage points|points|pp|pts)\b", t)
    if m:
        return "points", float(m.group(1))
    m = re.search(r"(\d+(?:\.\d+)?)\s*(?:%|percent|per cent)", t)
    if m:
        return "relative", min(float(m.group(1)) / 100, 1.0)
    for pat, f in _FRACTIONS:
        if re.search(pat, t):
            return "relative", f
    return None


def _band(a: dict, text: str) -> dict | None:
    """The analysis group the sentence names ("the low-usage band", "Blue card holders"), or None for everyone."""
    t = (text or "").lower()
    sigs = [s for s in (a.get("signals") or []) if s.get("group_n")]
    if not sigs:
        return None
    if re.search(r"\b(overall|everyone|all customers|across the board|every customer)\b", t):
        return None
    want_low = bool(re.search(r"\b(low|lowest|least|light|few|fewer|bottom)\b", t))
    want_high = bool(re.search(r"\b(high|highest|most|heavy|many|top)\b", t)) and not want_low
    stems: set[str] = set()
    for w in re.findall(r"[a-z]{3,}", t):
        stems.add(w[:5])
        for syn in _SYN.get(w, ()):
            stems.add(syn[:5])
    best, best_score = None, 0
    for i, s in enumerate(sigs):
        col = (s.get("column") or "").lower()
        lab = (s.get("label") or "").lower()
        hay = re.findall(r"[a-z]{3,}", (col.replace("_", " ") + " " + lab))
        score = sum(1 for w in hay if w[:5] in stems and w not in an_mod._STOP)
        if s.get("kind") == "categorical" and str(s.get("value") or "").lower() in t:
            score += 3
        if s.get("kind") == "numeric":
            if want_low and s.get("direction") == "low":
                score += 1
            if want_high and s.get("direction") == "high":
                score += 1
            if (want_low and s.get("direction") == "high") or (want_high and s.get("direction") == "low"):
                score -= 1
        score = score * 10 - i  # ties go to the stronger signal
        if score > best_score:
            best, best_score = s, score
    return best if best_score >= 10 else None


GROUP_PHRASE = re.compile(r"\b(?:among|amongst|for|in|within|with)\s+(?!the (?:data|table|next|coming|future)\b|a (?:year|month|week)\b|\d)(.{3,60}?)(?:[?.!,;]|$)", re.I)


def group_phrase(text: str) -> str | None:
    m = GROUP_PHRASE.search(text or "")
    return m.group(1).strip() if m else None


def band_from_frame(df, a: dict, text: str) -> dict | None:
    """A group named by a value in the table ("Blue card holders" -> Card_Category = Blue), its size and rate."""
    import pandas as pd
    tgt = a.get("target") or {}
    col_t, pos = tgt.get("column"), tgt.get("positive")
    if df is None or col_t not in getattr(df, "columns", []):
        return None
    y = df[col_t].astype(str) == str(pos)
    t = (text or "").lower()
    best = None
    for c in df.columns:
        if c == col_t or pd.api.types.is_numeric_dtype(df[c]):
            continue
        vals = df[c].dropna().astype(str)
        if vals.nunique() > 60:
            continue
        for v in vals.unique():
            vl = v.lower().strip()
            if len(vl) >= 3 and re.search(rf"(?<![a-z0-9]){re.escape(vl)}(?![a-z0-9])", t):
                mask = df[c].astype(str) == v
                n = int(mask.sum())
                if n >= 20 and (best is None or len(vl) > best[0]):
                    best = (len(vl), {"column": c, "label": an_mod.label_of(c), "kind": "categorical", "value": v,
                                      "group": f"{an_mod.label_of(c)} = {v}", "group_n": n, "group_rate": round(float(y[mask].mean()) * 100, 1)})
    return best[1] if best else None


def whatif(a: dict, text: str, band_override: dict | None = None) -> dict:
    """{ok, text, card} or {ok: False, text} when the sentence doesn't say how much, or names a group that isn't there."""
    ch = _change(text)
    b = a.get("baseline") or {}
    tgt = a.get("target") or {}
    noun = tgt.get("noun") or "the outcome"
    ent = (a.get("table") or {}).get("entities") or "rows"
    if not ch or not b.get("n"):
        return {"ok": False, "text": f"Say how much and for whom, for example “what if {noun} fell by a quarter among low-usage {ent}”. I work it out from this analysis."}
    kind, amt = ch
    band = band_override or _band(a, text)
    named = group_phrase(text)
    if band is None and named and not re.search(r"\b(overall|everyone|all)\b", named, re.I):
        groups = ", ".join(f"“{s['group']}”" for s in (a.get("signals") or [])[:3])
        return {"ok": False, "text": f"I couldn't find “{named}” in the analysis or the table, so I didn't guess. "
                                     f"Name a value from the data or one of the groups above ({groups})."}
    n_all, pos_all, rate_all = int(b["n"]), int(b["positives"]), float(b["rate"])
    if band:
        g_n, g_rate, g_label = int(band["group_n"]), float(band["group_rate"]), band["group"]
    else:
        g_n, g_rate, g_label = n_all, rate_all, f"all {n_all:,} {ent}"
    new_rate = max(0.0, g_rate * (1 - amt)) if kind == "relative" else max(0.0, g_rate - amt)
    before_pos = round(g_n * g_rate / 100)
    after_pos = round(g_n * new_rate / 100)
    saved = max(0, before_pos - after_pos)
    new_all = round((pos_all - saved) / n_all * 100, 1)
    how = (("by half" if abs(amt - 0.5) < 1e-9 else "by a quarter" if abs(amt - 0.25) < 1e-9 else "by a third" if abs(amt - 1 / 3) < 1e-9
            else f"by {round(amt * 100):g}%") if kind == "relative" else f"by {amt:g} points")
    if band and band.get("kind") == "categorical" and band.get("value") is not None:
        who = f"{ent} with {str(band.get('label') or '').lower()} {band['value']}"
    elif band:
        who = f"{ent} with {g_label[0].lower() + g_label[1:]}"
    else:
        who = g_label
    verb = "fewer would leave" if noun in ("churn", "attrition") else f"fewer would have the outcome"
    text_out = (f"If {noun} among {who} ({g_n:,} {ent}, {g_rate:g}% today) fell {how} to {round(new_rate, 1):g}%, about {saved:,} {verb}, "
                f"and overall {noun} would go from {rate_all:g}% to {new_all:g}%.") if band else \
               (f"If {noun} fell {how} across all {n_all:,} {ent}, from {rate_all:g}% to {round(new_rate, 1):g}%, about {saved:,} {verb}.")
    assumptions = [
        (f"The change applies only to {who}; every other {ent[:-1] if ent.endswith('s') else ent} keeps its past rate." if band
         else f"The change applies evenly to every {ent[:-1] if ent.endswith('s') else ent}."),
        (f"“{how[3:]}” means {round(amt * 100):g}% of that group's rate (relative), not {round(amt * 100):g} points." if kind == "relative"
         else f"{amt:g} points means the rate itself drops by {amt:g} percentage points."),
        f"The rates are from the period in this table; this is arithmetic on them, not a forecast.",
        f"The same {n_all:,} {ent}: no new {ent}, no changes to price or product.",
    ]
    card = {"group": g_label if band else f"All {ent}", "group_n": g_n, "column": (band or {}).get("column"),
            "before": {"group_rate": g_rate, "group_count": before_pos, "overall_rate": rate_all, "overall_count": pos_all},
            "after": {"group_rate": round(new_rate, 1), "group_count": after_pos, "overall_rate": new_all, "overall_count": pos_all - saved},
            "change": {"kind": kind, "amount": amt, "label": how}, "saved": saved, "noun": noun, "entities": ent, "assumptions": assumptions}
    return {"ok": True, "text": text_out, "card": card}
