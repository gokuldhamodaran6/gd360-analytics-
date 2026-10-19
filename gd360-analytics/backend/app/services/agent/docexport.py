"""2026-10-15 (R51): a block document as Word, PDF and Markdown - the same blocks the page shows.

Word (python-docx): the title, headings, paragraphs with their bold figures and source notes, figures as charts
drawn to scale (PNG, Pillow), tables with a header row, the limitation as a shaded note, sources at the end.
PDF (reportlab): the same, with figures drawn as vector bars. Markdown: for pasting anywhere.
"""
from __future__ import annotations

import io
import re

from PIL import Image, ImageDraw, ImageFont

from .docblocks import plain

INK = (11, 15, 14)
MUTED = (94, 106, 102)
LINE = (221, 226, 224)
TONE = {"bad": (194, 65, 12), "good": (10, 124, 90), "neutral": (154, 165, 160), "main": (47, 111, 228), "top": (31, 85, 192)}


def nice_axis(vmax: float) -> tuple[float, float]:
    """R57: an axis that fits any scale - (top, step) with 4-6 round steps (0, 2M, 4M, 6M, 8M; never 799k ticks)."""
    import math
    if not vmax or vmax <= 0:
        return 1.0, 0.25
    raw = vmax / 5
    mag = 10 ** math.floor(math.log10(raw))
    step = next(m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw)
    return math.ceil(vmax / step - 1e-9) * step, step


def tick_text(v: float, unit: str = "") -> str:
    a = abs(v)
    if a >= 1e9:
        body = f"{v / 1e9:g}B"
    elif a >= 1e6:
        body = f"{v / 1e6:g}M"
    elif a >= 1e4:
        body = f"{v / 1e3:g}K"
    else:
        body = f"{v:g}"
    return body + (unit or "")


def bar_text(b: dict, unit: str = "") -> str:
    return str(b.get("display") or tick_text(float(b["value"]), unit))


def _ticks(top: float, step: float) -> list[float]:
    out, k = [], 0.0
    while k <= top + 1e-9 and len(out) < 12:
        out.append(round(k, 10))
        k += step
    return out


def _src_text(doc: dict, cites: list[str]) -> str:
    by = {s["id"]: s for s in doc.get("sources") or []}
    labels = [by[c]["label"] for c in cites or [] if c in by]
    return " · ".join(labels)


def _runs(text: str) -> list[tuple[str, bool]]:
    """'a **16.3%** rate' -> [('a ', False), ('16.3%', True), (' rate', False)]"""
    out = []
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text or "")):
        if part:
            out.append((part, i % 2 == 1))
    return out


def _font(size: int, bold: bool = False):
    try:
        return ImageFont.load_default(size=size)
    except TypeError:  # an older Pillow
        return ImageFont.load_default()


def line_png(fig: dict, width: int = 1400) -> bytes:
    """R57: a series over time as a line, the first and last points labelled."""
    pts = fig.get("bars") or []
    h, left, right, top, bottom = 560, 130, 60, 40, 80
    im = Image.new("RGB", (width, h), (255, 255, 255))
    d = ImageDraw.Draw(im)
    f_ax, f_val = _font(22), _font(24, True)
    vals = [float(p["value"]) for p in pts]
    lo = min(0.0, min(vals))
    top_v, step = nice_axis(max(vals) - lo if lo < 0 else max(vals))
    pw, ph = width - left - right, h - top - bottom

    def xy(i: int, v: float) -> tuple[float, float]:
        x = left + (pw * i / max(1, len(pts) - 1))
        return x, top + ph * (1 - (v - lo) / (top_v or 1))

    for k in _ticks(top_v, step):
        y = xy(0, k + lo)[1]
        d.line([(left, y), (left + pw, y)], fill=(240, 242, 241), width=2)
        d.text((left - 14, y), tick_text(k + lo, fig.get("unit") or ""), fill=MUTED, font=f_ax, anchor="rm")
    d.line([xy(i, v) for i, v in enumerate(vals)], fill=TONE["main"], width=5, joint="curve")
    every = max(1, len(pts) // 8)
    for i, p in enumerate(pts):
        if i % every == 0 or i == len(pts) - 1:
            d.text((xy(i, 0)[0], top + ph + 18), str(p["label"])[:10], fill=MUTED, font=f_ax, anchor="ma")
    for i in (0, len(pts) - 1):
        x, y = xy(i, vals[i])
        d.ellipse([x - 7, y - 7, x + 7, y + 7], fill=TONE["main"])
        d.text((x + (12 if i == 0 else -12), y - 14), bar_text(pts[i], fig.get("unit") or ""), fill=INK, font=f_val, anchor="ld" if i == 0 else "rd")
    buf = io.BytesIO()
    im.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def chart_png(fig: dict, width: int = 1400) -> bytes:
    """A horizontal bar chart to scale, with the baseline as a dashed line and each value labelled."""
    if fig.get("kind") == "line":
        return line_png(fig, width)
    bars = fig.get("bars") or []
    row_h, top, left, right = 64, 70, 430, 190
    h = top + row_h * len(bars) + 70
    im = Image.new("RGB", (width, h), (255, 255, 255))
    d = ImageDraw.Draw(im)
    f_lab, f_val, f_ax = _font(26), _font(26, True), _font(22)
    vmax, step = nice_axis(max([b["value"] for b in bars] + [(fig.get("baseline") or {}).get("value") or 0, 1e-9]))
    plot_w = width - left - right
    x0 = left

    def xs(v: float) -> float:
        return x0 + plot_w * (max(0.0, v) / vmax)

    for k in _ticks(vmax, step):
        d.line([(xs(k), top - 10), (xs(k), top + row_h * len(bars))], fill=(240, 242, 241), width=2)
        d.text((xs(k), top + row_h * len(bars) + 14), tick_text(k, fig.get("unit") or ""), fill=MUTED, font=f_ax, anchor="ma")
    for i, b in enumerate(bars):
        y = top + i * row_h + 14
        d.text((left - 18, y + 17), str(b["label"])[:32], fill=INK, font=f_lab, anchor="rm")
        d.rounded_rectangle([x0, y, max(x0 + 4, xs(b["value"])), y + 36], radius=8, fill=TONE.get(b.get("tone"), TONE["neutral"]))
        d.text((xs(b["value"]) + 12, y + 18), bar_text(b, fig.get("unit") or ""), fill=INK, font=f_val, anchor="lm")
    base = fig.get("baseline") or {}
    if base.get("value") is not None:
        bx = xs(base["value"])
        yy = top - 14
        while yy < top + row_h * len(bars):
            d.line([(bx, yy), (bx, min(yy + 10, top + row_h * len(bars)))], fill=INK, width=3)
            yy += 18
        d.text((bx + 10, top - 40), base.get("label") or "", fill=INK, font=f_ax)
    d.line([(x0, top + row_h * len(bars)), (x0 + plot_w, top + row_h * len(bars))], fill=LINE, width=2)
    buf = io.BytesIO()
    im.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


# ---- Markdown -------------------------------------------------------------------------------

def markdown(doc: dict) -> str:
    out = []
    for b in doc["blocks"]:
        t = b["type"]
        src = _src_text(doc, b.get("cites") or [])
        if t == "title":
            out += [f"# {plain(b['text'])}", ""]
        elif t == "byline":
            out += [f"_{b['text']}_", ""]
        elif t == "heading":
            out += [f"## {plain(b['text'])}", ""]
        elif t == "paragraph":
            out += [b["text"] + (f" _({src})_" if src else ""), ""]
        elif t == "callout":
            out += [f"> **{b.get('label') or 'Note'}.** {plain(b['text'])}", ""]
        elif t == "figure":
            out += [f"**{b['title']}** · {b.get('subtitle') or ''}", "", "| Group | Value |", "|---|---:|"]
            out += [f"| {x['label']} | {bar_text(x, b.get('unit') or '')} |" for x in b.get("bars") or []]
            out += [""] + ([f"_{src}_", ""] if src else [])
        elif t == "table":
            cols = b.get("columns") or []
            out += ["| " + " | ".join(c["label"] for c in cols) + " |", "|" + "|".join("---:" if c.get("align") == "right" else "---" for c in cols) + "|"]
            out += ["| " + " | ".join(str(row.get(c["key"], "")) for c in cols) + " |" for row in b.get("rows") or []]
            out += ([f"_{b['note']}_"] if b.get("note") else []) + [""]
        elif t == "list":
            out += [f"- {x}" for x in b.get("items") or []] + [""]
        elif t == "kpis":
            items = b.get("items") or []
            out += ["| " + " | ".join(str(k.get("label") or "") for k in items) + " |", "|" + "|".join("---:" for _ in items) + "|",
                    "| " + " | ".join(f"**{k.get('value') or ''}**" for k in items) + " |", ""]
        elif t == "sources":
            srcs = doc.get("sources") or []
            if srcs:
                out += ["---", "", "**Sources**", ""] + [f"{i + 1}. {s['label']}" + (f" ({s['url']})" if s.get("url") else "") for i, s in enumerate(srcs)] + [""]
    out.append(f"_Every figure re-checked against the data · {doc.get('checked', 0)} checked_")
    return "\n".join(out)


# ---- Word -----------------------------------------------------------------------------------

def docx_bytes(doc: dict) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.0)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.font.color.rgb = RGBColor(0x2A, 0x33, 0x30)

    def shade(par, hexcolor: str):
        p_pr = par._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexcolor)
        p_pr.append(shd)

    def add_text(par, text: str, size: float | None = None, color: tuple | None = None):
        for part, bold in _runs(text):
            r = par.add_run(part)
            r.bold = bold
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor(*color)

    def note(text: str):
        if not text:
            return
        p = d.add_paragraph()
        r = p.add_run(f"Source: {text}")
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(*MUTED)

    for b in doc["blocks"]:
        t = b["type"]
        if t == "title":
            p = d.add_paragraph()
            r = p.add_run(plain(b["text"]))
            r.bold = True
            r.font.size = Pt(26)
            r.font.color.rgb = RGBColor(*INK)
        elif t == "byline":
            p = d.add_paragraph()
            r = p.add_run(b["text"])
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(*MUTED)
        elif t == "heading":
            h = d.add_heading(plain(b["text"]), level=1)
            for r in h.runs:
                r.font.color.rgb = RGBColor(*INK)
                r.font.size = Pt(15)
        elif t == "paragraph":
            p = d.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            add_text(p, b["text"])
            note(_src_text(doc, b.get("cites") or []))
        elif t == "callout":
            p = d.add_paragraph()
            shade(p, "FFF4DF")
            r = p.add_run(f"{b.get('label') or 'Note'}. ")
            r.bold = True
            r.font.color.rgb = RGBColor(0x8A, 0x5D, 0x00)
            r2 = p.add_run(plain(b["text"]))
            r2.font.color.rgb = RGBColor(0x8A, 0x5D, 0x00)
        elif t == "list":
            for x in b.get("items") or []:
                p = d.add_paragraph(style="List Bullet")
                add_text(p, str(x))
        elif t == "kpis":
            items = b.get("items") or []
            if items:
                tb = d.add_table(rows=2, cols=len(items))
                for j, k in enumerate(items):
                    c0 = tb.rows[0].cells[j]
                    c0.text = ""
                    r0 = c0.paragraphs[0].add_run(str(k.get("label") or ""))
                    r0.font.size = Pt(9)
                    r0.font.color.rgb = RGBColor(*MUTED)
                    c1 = tb.rows[1].cells[j]
                    c1.text = ""
                    r1 = c1.paragraphs[0].add_run(str(k.get("value") or ""))
                    r1.bold = True
                    r1.font.size = Pt(16)
                d.add_paragraph()
        elif t == "figure":
            p = d.add_paragraph()
            r = p.add_run(b["title"])
            r.bold = True
            if b.get("subtitle"):
                p2 = d.add_paragraph()
                r2 = p2.add_run(b["subtitle"])
                r2.font.size = Pt(9)
                r2.font.color.rgb = RGBColor(*MUTED)
            d.add_picture(io.BytesIO(chart_png(b)), width=Cm(16))
            d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
            note(_src_text(doc, b.get("cites") or []))
        elif t == "table":
            cols = b.get("columns") or []
            rows = b.get("rows") or []
            tb = d.add_table(rows=1 + len(rows), cols=len(cols))
            tb.alignment = WD_TABLE_ALIGNMENT.LEFT

            def bottom(cell, color: str, size: int):
                tc_pr = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                el = OxmlElement("w:bottom")
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), str(size))
                el.set(qn("w:color"), color)
                borders.append(el)
                tc_pr.append(borders)

            for ri, row_ in enumerate(tb.rows):
                for cell in row_.cells:
                    bottom(cell, "C9D1CE" if ri == 0 else "E6E9E8", 8 if ri == 0 else 4)
            for j, c in enumerate(cols):
                cell = tb.rows[0].cells[j]
                cell.text = ""
                rr = cell.paragraphs[0].add_run(c["label"])
                rr.bold = True
                rr.font.size = Pt(9.5)
                rr.font.color.rgb = RGBColor(*MUTED)
                if c.get("align") == "right":
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for i, row in enumerate(rows):
                for j, c in enumerate(cols):
                    cell = tb.rows[i + 1].cells[j]
                    cell.text = ""
                    rr = cell.paragraphs[0].add_run(str(row.get(c["key"], "")))
                    rr.font.size = Pt(9.5)
                    rr.bold = False
                    rr.font.color.rgb = RGBColor(0x2A, 0x33, 0x30)
                    if c.get("align") == "right":
                        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            note(_src_text(doc, b.get("cites") or []))
        elif t == "sources":
            srcs = doc.get("sources") or []
            if srcs:
                h = d.add_heading("Sources", level=2)
                for r in h.runs:
                    r.font.color.rgb = RGBColor(*INK)
                for i, s in enumerate(srcs):
                    p = d.add_paragraph()
                    r = p.add_run(f"{i + 1}. {s['label']}" + (f" - {s['url']}" if s.get("url") else ""))
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(*MUTED)
    p = d.add_paragraph()
    r = p.add_run(f"Every figure re-checked against the data before export · {doc.get('checked', 0)} checked · GD360")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(*MUTED)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ---- PDF ------------------------------------------------------------------------------------

def pdf_bytes(doc: dict) -> bytes:
    from reportlab.graphics.shapes import Drawing, Line, Rect, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ink = colors.Color(*(c / 255 for c in INK))
    muted = colors.Color(*(c / 255 for c in MUTED))
    sty = {
        "title": ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=26, leading=30, textColor=ink, spaceAfter=8),
        "byline": ParagraphStyle("b", fontName="Helvetica", fontSize=9, leading=12, textColor=muted, spaceAfter=14),
        "h2": ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=ink, spaceBefore=14, spaceAfter=6),
        "p": ParagraphStyle("p", fontName="Helvetica", fontSize=10.5, leading=15.5, textColor=colors.Color(42 / 255, 51 / 255, 48 / 255), spaceAfter=3),
        "src": ParagraphStyle("s", fontName="Helvetica", fontSize=8, leading=10, textColor=muted, spaceAfter=8),
        "call": ParagraphStyle("c", fontName="Helvetica", fontSize=10, leading=14.5, textColor=colors.Color(138 / 255, 93 / 255, 0),
                               backColor=colors.Color(1, 244 / 255, 223 / 255), borderPadding=8, spaceBefore=6, spaceAfter=12),
        "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=9, leading=12),
        "cellh": ParagraphStyle("cellh", fontName="Helvetica-Bold", fontSize=9, leading=12, textColor=muted),
    }

    def esc(s: str) -> str:
        return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def rich(s: str) -> str:
        return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", esc(s))

    W = A4[0] - 4.4 * cm
    story = []
    for b in doc["blocks"]:
        t = b["type"]
        src = _src_text(doc, b.get("cites") or [])
        if t == "title":
            story.append(Paragraph(esc(plain(b["text"])), sty["title"]))
        elif t == "byline":
            story.append(Paragraph(esc(b["text"]), sty["byline"]))
        elif t == "heading":
            story.append(Paragraph(esc(plain(b["text"])), sty["h2"]))
        elif t == "paragraph":
            story.append(Paragraph(rich(b["text"]), sty["p"]))
            if src:
                story.append(Paragraph(esc("Source: " + src), sty["src"]))
        elif t == "callout":
            story.append(Paragraph(f"<b>{esc(b.get('label') or 'Note')}.</b> {esc(plain(b['text']))}", sty["call"]))
        elif t == "list":
            for x in b.get("items") or []:
                story.append(Paragraph(rich(str(x)), sty["p"], bulletText="•"))
            story.append(Spacer(1, 4))
        elif t == "kpis":
            items = b.get("items") or []
            if items:
                data = [[Paragraph(esc(str(k.get("label") or "")), sty["cellh"]) for k in items],
                        [Paragraph(f"<b>{esc(str(k.get('value') or ''))}</b>", ParagraphStyle("kv", fontName="Helvetica-Bold", fontSize=15, leading=19, textColor=ink)) for k in items]]
                tb = Table(data, colWidths=[W / len(items)] * len(items))
                tb.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.5, colors.Color(*(v / 255 for v in LINE))),
                                        ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.Color(*(v / 255 for v in LINE))),
                                        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
                story.append(tb)
                story.append(Spacer(1, 8))
        elif t == "figure" and b.get("kind") == "line":
            pts = b.get("bars") or []
            h, left, top, bot = 170, 46, 14, 26
            dw = Drawing(W, h)
            vals = [float(x["value"]) for x in pts] or [0.0]
            lo = min(0.0, min(vals))
            vmax, step = nice_axis(max(vals) - lo if lo < 0 else max(vals))
            pw, ph = W - left - 12, h - top - bot

            def lxy(i: int, v: float) -> tuple[float, float]:
                return left + pw * i / max(1, len(pts) - 1), bot + ph * ((v - lo) / (vmax or 1))

            for k in _ticks(vmax, step):
                y = lxy(0, k + lo)[1]
                dw.add(Line(left, y, left + pw, y, strokeColor=colors.Color(240 / 255, 242 / 255, 241 / 255), strokeWidth=0.6))
                dw.add(String(left - 4, y - 2.5, tick_text(k + lo, b.get("unit") or ""), fontName="Helvetica", fontSize=7, fillColor=muted, textAnchor="end"))
            main = colors.Color(*(v / 255 for v in TONE["main"]))
            for i in range(len(pts) - 1):
                x1, y1 = lxy(i, vals[i])
                x2, y2 = lxy(i + 1, vals[i + 1])
                dw.add(Line(x1, y1, x2, y2, strokeColor=main, strokeWidth=1.8))
            every = max(1, len(pts) // 8)
            for i, x in enumerate(pts):
                if i % every == 0 or i == len(pts) - 1:
                    dw.add(String(lxy(i, 0)[0], bot - 14, str(x["label"])[:10], fontName="Helvetica", fontSize=7, fillColor=muted, textAnchor="middle"))
            parts = [Paragraph(f"<b>{esc(b['title'])}</b>", sty["p"])]
            if b.get("subtitle"):
                parts.append(Paragraph(esc(b["subtitle"]), sty["src"]))
            parts.append(dw)
            if src:
                parts.append(Paragraph(esc("Source: " + src), sty["src"]))
            story.append(KeepTogether(parts))
        elif t == "figure":
            bars = b.get("bars") or []
            row_h, top, left = 22, 26, 150
            h = top + row_h * len(bars) + 26
            dw = Drawing(W, h)
            vmax, step = nice_axis(max([x["value"] for x in bars] + [(b.get("baseline") or {}).get("value") or 0, 1e-9]))
            pw = W - left - 60

            def xs(v: float) -> float:
                return left + pw * (max(0.0, v) / vmax)

            ybot = 20
            for k in _ticks(vmax, step):
                dw.add(String(xs(k), ybot - 12, tick_text(k, b.get("unit") or ""), fontName="Helvetica", fontSize=7, fillColor=muted, textAnchor="middle"))
            for i, x in enumerate(bars):
                y = h - top - (i + 1) * row_h + 4
                c = TONE.get(x.get("tone"), TONE["neutral"])
                dw.add(String(left - 6, y + 4, str(x["label"])[:34], fontName="Helvetica", fontSize=8, fillColor=ink, textAnchor="end"))
                dw.add(Rect(left, y, max(2, xs(x["value"]) - left), 13, fillColor=colors.Color(*(v / 255 for v in c)), strokeColor=None, rx=3, ry=3))
                dw.add(String(xs(x["value"]) + 4, y + 3.5, bar_text(x, b.get("unit") or ""), fontName="Helvetica-Bold", fontSize=8, fillColor=ink))
            base = b.get("baseline") or {}
            if base.get("value") is not None:
                bx = xs(base["value"])
                dw.add(Line(bx, ybot, bx, h - 14, strokeColor=ink, strokeWidth=1, strokeDashArray=[3, 3]))
                dw.add(String(bx + 4, h - 12, base.get("label") or "", fontName="Helvetica", fontSize=7.5, fillColor=ink))
            dw.add(Line(left, ybot, left + pw, ybot, strokeColor=colors.Color(*(v / 255 for v in LINE)), strokeWidth=1))
            parts = [Paragraph(f"<b>{esc(b['title'])}</b>", sty["p"])]
            if b.get("subtitle"):
                parts.append(Paragraph(esc(b["subtitle"]), sty["src"]))
            parts.append(dw)
            if src:
                parts.append(Paragraph(esc("Source: " + src), sty["src"]))
            story.append(KeepTogether(parts))
        elif t == "table":
            cols = b.get("columns") or []
            data = [[Paragraph(esc(c["label"]), sty["cellh"]) for c in cols]]
            for row in b.get("rows") or []:
                data.append([Paragraph(esc(str(row.get(c["key"], ""))), sty["cell"]) for c in cols])
            tb = Table(data, colWidths=[W / max(1, len(cols))] * len(cols), repeatRows=1)
            style = [("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.Color(*(v / 255 for v in LINE))),
                     ("LINEBELOW", (0, 1), (-1, -1), 0.4, colors.Color(236 / 255, 238 / 255, 237 / 255)),
                     ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]
            for j, c in enumerate(cols):
                if c.get("align") == "right":
                    style.append(("ALIGN", (j, 0), (j, -1), "RIGHT"))
            tb.setStyle(TableStyle(style))
            story.append(tb)
            story.append(Spacer(1, 6))
            if b.get("note"):
                story.append(Paragraph(esc(b["note"]), sty["src"]))
            if src:
                story.append(Paragraph(esc("Source: " + src), sty["src"]))
        elif t == "sources":
            srcs = doc.get("sources") or []
            if srcs:
                story.append(Paragraph("Sources", sty["h2"]))
                for i, s in enumerate(srcs):
                    story.append(Paragraph(esc(f"{i + 1}. {s['label']}" + (f" - {s['url']}" if s.get("url") else "")), sty["src"]))
    story.append(Spacer(1, 10))
    story.append(Paragraph(esc(f"Every figure re-checked against the data before export · {doc.get('checked', 0)} checked · GD360"), sty["src"]))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
                      title=doc.get("title") or "Document", author="GD360").build(story)
    return buf.getvalue()
