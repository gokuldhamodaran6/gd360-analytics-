"""2026-10-15 (R49): designed dashboards ("boards") - a dashboard is designed, not generated.

A board is a Dashboard row with layout_version 4 and a JSON `board_spec`:

    {version, title, subtitle, theme, source{source_id, table, name}, entities, outcome?{column, positive, verb, noun},
     dates?{column}, filters[{key, label, column, values}], pages[{id, name, blurb, blocks[{id, kind, span, title, ...}]}],
     definitions[str], about[str], model?{ap, auc, test_n, ids, scores, importance}, signals[...], built_at}

`design()` writes the spec from the table (and the conversation's analysis when there is one): a Home page,
an Overview, then pages that fit the data - Signals / Segments / Customers for a yes/no outcome, Breakdown /
Trends for measures over time. `compute()` runs one page live on the source with the chosen filters and returns
plain numbers for frontend/src/board/BoardView.tsx to draw; `apply_edit()` changes the spec by talking; and
`workbook()` / `deck()` export the computed pages, so the Excel and PowerPoint carry the same pages and theme.

Numbers are computed here, never by a model. Every block carries its definition, and the Home page lists them.
"""
from __future__ import annotations

import copy
import io
import math
import re
import threading
import time
from datetime import datetime

import numpy as np
import pandas as pd
from sqlalchemy.orm import Session
from sqlalchemy.orm.attributes import flag_modified

from ... import models
from .. import ml_studio
from . import analysis

MAX_ROWS = 150_000
SCORES_MAX = 60_000

THEMES = {
    "ink": {"label": "Ink", "mode": "light", "accent": "#2F5BFF", "warn": "#F26B3A", "good": "#15803D", "bg": "#F5F6F8", "card": "#FFFFFF", "ink": "#0B0F14"},
    "forest": {"label": "Forest", "mode": "dark", "accent": "#8FD3B4", "warn": "#F4A261", "good": "#8FD3B4", "bg": "#0F2E22", "card": "#163A2C", "ink": "#F2F7F4"},
    "navy": {"label": "Navy", "mode": "dark", "accent": "#7FB3F5", "warn": "#F4A261", "good": "#7FD4A8", "bg": "#0E1B33", "card": "#16264A", "ink": "#EEF3FB"},
    "graphite": {"label": "Graphite", "mode": "dark", "accent": "#A5B4FC", "warn": "#FB923C", "good": "#4ADE80", "bg": "#0F1115", "card": "#181B21", "ink": "#F3F4F6"},
    "sand": {"label": "Sand", "mode": "light", "accent": "#B45309", "warn": "#B91C1C", "good": "#15803D", "bg": "#F7F4EE", "card": "#FFFFFF", "ink": "#1C1917"},
}
THEME_WORDS = {"ink": r"\b(ink|light|white|default)\b", "forest": r"\b(forest|green)\b", "navy": r"\b(navy|blue|dark blue)\b",
               "graphite": r"\b(graphite|dark|black|night)\b", "sand": r"\b(sand|warm|beige|cream)\b"}

MONEY = re.compile(r"amt|amount|revenue|sales|spend|price|cost|profit|margin|limit|bal(ance)?\b|_bal|value|fee|income_amt|adr|gmv|payment|discount|tax|tip|wage|salary|payroll|cogs|refund", re.I)
TENURE = re.compile(r"months?_on|tenure|months_on_book|age_days|since|lead_time", re.I)
DIM_PRIORITY = ("category", "segment", "plan", "tier", "type", "product", "region", "channel", "market", "country", "city", "store", "hotel",
                "income", "education", "status", "team", "department", "gender", "marital")


# ---- json safety ---------------------------------------------------------------------------

def _f(v, nd: int = 4):
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(x):
        return None
    return round(x, nd)


def _clean(o):
    if isinstance(o, dict):
        return {k: _clean(v) for k, v in o.items() if not str(k).startswith("_")}
    if isinstance(o, (list, tuple)):
        return [_clean(v) for v in o]
    if isinstance(o, (np.integer,)):
        return int(o)
    if isinstance(o, (np.floating, float)):
        return _f(o)
    if isinstance(o, (np.bool_,)):
        return bool(o)
    return o


# ---- frames, cached briefly per person ------------------------------------------------------

_CACHE: dict[tuple, tuple[float, pd.DataFrame, bool]] = {}
_LOCK = threading.Lock()


def frame(db: Session, user: models.User, source_id: str, table: str | None) -> tuple[pd.DataFrame, bool]:
    key = (user.id, source_id, table)
    with _LOCK:
        hit = _CACHE.get(key)
        if hit and time.time() - hit[0] < 300:
            return hit[1], hit[2]
    df, _entry, sampled = ml_studio.load_frame(db, user, source_id, table, cap=MAX_ROWS)
    df = df.reset_index(drop=True)
    with _LOCK:
        if len(_CACHE) > 6:
            _CACHE.pop(next(iter(_CACHE)))
        _CACHE[key] = (time.time(), df, sampled)
    return df, sampled


def forget(source_id: str) -> None:
    with _LOCK:
        for k in [k for k in _CACHE if k[1] == source_id]:
            _CACHE.pop(k, None)


# ---- reading the table ------------------------------------------------------------------------

def _date_column(df: pd.DataFrame) -> str | None:
    best, best_n = None, 0
    for c in df.columns:
        s = df[c]
        if pd.api.types.is_datetime64_any_dtype(s):
            n = s.notna().sum()
        elif s.dtype == object and analysis.DATE_NAMES.search(str(c)):
            try:
                p = pd.to_datetime(s.dropna().astype(str).head(300), errors="coerce", format="mixed")
            except (TypeError, ValueError):
                p = pd.to_datetime(s.dropna().astype(str).head(300), errors="coerce")
            if p.notna().mean() < 0.85:
                continue
            n = s.notna().sum()
        else:
            continue
        if n > best_n:
            best, best_n = c, n
    return best


def _dates(df: pd.DataFrame, col: str) -> pd.Series:
    s = df[col]
    if pd.api.types.is_datetime64_any_dtype(s):
        return s
    try:
        return pd.to_datetime(s, errors="coerce", format="mixed")
    except (TypeError, ValueError):
        return pd.to_datetime(s, errors="coerce")


def profile(df: pd.DataFrame) -> dict:
    """{ids, measures, dims, date, outcome?}"""
    n = len(df)
    ids, measures, dims = [], [], []
    for c in df.columns:
        s = df[c]
        nun = s.nunique(dropna=True)
        if nun <= 1:
            continue
        if n > 50 and nun >= 0.9 * n and not pd.api.types.is_float_dtype(s):
            ids.append(c)
            continue
        if analysis.ID_NAMES.search(str(c)) and nun > 0.5 * n:
            ids.append(c)
            continue
        if pd.api.types.is_bool_dtype(s):
            continue
        if pd.api.types.is_numeric_dtype(s):
            if re.search(r"naive_bayes|classifier", str(c), re.I):
                continue
            if nun > 10 or pd.api.types.is_float_dtype(s):
                measures.append(c)
            continue
        if 2 <= nun <= 30 and not analysis.DATE_NAMES.search(str(c)):
            dims.append(c)
    dims.sort(key=lambda c: (next((i for i, w in enumerate(DIM_PRIORITY) if w in str(c).lower()), len(DIM_PRIORITY)), df[c].nunique()))
    measures.sort(key=lambda c: (0 if MONEY.search(str(c)) else 1))
    out = {"ids": ids, "measures": measures, "dims": dims, "date": _date_column(df)}
    cands = analysis._binary_candidates(df)
    named = [c for c, sc in cands if sc >= 2.0]
    if named:
        col, pos = analysis.pick_target(df, named[0], "")
        out["outcome"] = {"column": col, "positive": pos}
    return out


def short(col: str) -> str:
    """Total_Trans_Ct -> "transaction count": the label without the words every column of a kind repeats."""
    lab = analysis.label_of(col)
    lab = re.sub(r"^(total|avg|average|sum of|number of)\s+", "", lab)
    lab = re.sub(r"\s+in 12 months$", "", lab) if len(lab) > 22 else lab
    return lab or analysis.label_of(col)


def _cap(s: str) -> str:
    return s[:1].upper() + s[1:] if s else s


def _relabel(s: dict) -> dict:
    """A signal's group names in short words: "transaction count under 54" against "75 or more"."""
    out = dict(s)
    lab = short(s["column"])
    out["label"] = lab
    if s.get("kind") == "numeric" and len(s.get("threshold") or []) == 2:
        q1, q2 = s["threshold"]
        integer = float(q1).is_integer() and float(q2).is_integer()
        lo = f"{lab} under {analysis._fmt_num(q1, integer)}" if q1 != q2 else f"{lab} of {analysis._fmt_num(q1, integer)} or less"
        hi = f"{lab} of {analysis._fmt_num(q2, integer)} or more" if q1 != q2 else f"{lab} above {analysis._fmt_num(q1, integer)}"
        out["group"], out["compare"] = (lo, hi) if s.get("direction") == "low" else (hi, lo)
    elif s.get("kind") == "categorical":
        out["group"], out["compare"] = f"{lab}: {s.get('value')}", f"{lab}: {s.get('compare_value')}"
    return out


def _unit(col: str, df: pd.DataFrame | None = None) -> str:
    if MONEY.search(str(col)) and not re.search(r"ratio|chng|change|(^|_)ct\b|(^|_)count\b|(^|_)cnt\b", str(col), re.I):
        return "money"
    if re.search(r"ratio|rate|pct|percent|share|chng|change", str(col), re.I):
        return "ratio"
    return "num"


# ---- design -----------------------------------------------------------------------------------

def _title_for(outcome: dict | None, entities: str, source_name: str, measures: list[str]) -> tuple[str, str]:
    if outcome:
        noun = outcome["noun"]
        if noun == "churn":
            return "Retention pulse", f"Who churns, why, and which {entities} are at risk now"
        if noun == "cancellation":
            return "Cancellation pulse", f"Which {entities} cancel, why, and which are at risk now"
        return f"{noun[:1].upper() + noun[1:]} pulse", f"What drives {noun}, and where the risk is now"
    if measures:
        m = analysis.label_of(measures[0])
        return f"{m[:1].upper() + m[1:]} overview", f"{source_name}: the main numbers, where they come from, and how they move"
    return f"{source_name} overview", "The main numbers in this table"


def design(db: Session, user: models.User, source_id: str, table: str | None, *, question: str = "", analysis_result: dict | None = None,
           theme: str = "ink", name: str | None = None) -> dict:
    tables = ml_studio.source_tables(db, user, [source_id])
    if not tables:
        raise analysis.AnalysisError("Connect a table first; there is nothing to build a dashboard from.")
    entry = next((t for t in tables if t["table"] == table), None) if table else None
    entry = entry or analysis.pick_table(db, user, [source_id], question)
    df, sampled = frame(db, user, entry["source_id"], entry["table"])
    if len(df) < 20:
        raise analysis.AnalysisError(f"{entry['table']} has only {len(df)} rows; a dashboard needs more to be useful.")
    prof = profile(df)
    entities = analysis._entities([str(c) for c in df.columns], entry["table"])
    if not prof.get("outcome") and not (analysis_result and analysis_result.get("target")):
        # a table of measures: a row is only an "order" or a "customer" when the table says so
        t = (entry["table"] or "").lower()
        entities = next((pl for w, pl in (("order", "orders"), ("booking", "bookings"), ("customer", "customers"), ("transaction", "transactions"),
                                         ("invoice", "invoices"), ("ticket", "tickets"), ("employee", "employees"), ("session", "sessions")) if w in t), "rows")
    spec: dict = {"version": 1, "theme": theme if theme in THEMES else "ink", "entities": entities,
                  "source": {"source_id": entry["source_id"], "table": entry["table"], "name": entry["source"]},
                  "rows_at_build": int(len(df)), "sampled": bool(sampled), "built_at": datetime.utcnow().isoformat() + "Z",
                  "question": question[:300]}
    dims = prof["dims"][:6]
    if prof["date"]:
        spec["dates"] = {"column": prof["date"]}
    outcome = prof.get("outcome")
    if analysis_result and analysis_result.get("target", {}).get("column") in df.columns:
        outcome = {"column": analysis_result["target"]["column"], "positive": analysis_result["target"]["positive"]}
    if outcome:
        r = analysis.analyse(df, entry, sampled, question or "why do rows have this outcome", {"recipe": "outcome", "target": outcome["column"]}, keep_scores=True)
        verb, noun = r["target"]["verb"], r["target"]["noun"]
        spec["outcome"] = {"column": r["target"]["column"], "positive": r["target"]["positive"], "verb": verb, "noun": noun}
        spec["signals"] = [_relabel({k: s.get(k) for k in ("column", "label", "kind", "direction", "group", "compare", "threshold", "value", "compare_value", "importance")})
                           for s in r["signals"][:8]]
        m = r.get("model")
        if m:
            spec["model"] = {"ap": m["ap"], "auc": m["auc"], "test_n": m["test_n"], "baseline_ap": m["baseline_ap"],
                             "importance": [d | {"label": short(d["column"])} for d in (r.get("importance") or [])], "leaks": [l["column"] for l in m.get("leaks") or []]}
            sc = r.get("_scores")
            id_col = r.get("id_column")
            if sc is not None and len(df) <= SCORES_MAX:
                spec["model"]["scores"] = [round(float(x), 3) for x in sc]
                spec["model"]["id_column"] = id_col
                spec["model"]["rows"] = int(len(df))
        title, subtitle = _title_for(spec["outcome"], entities, entry["source"], prof["measures"])
        spec["title"], spec["subtitle"] = name or title, subtitle
        spec["pages"] = _outcome_pages(df, spec, prof, dims)
        spec["definitions"] = [
            f"{noun} rate = {entities} with {r['target']['column']} = “{r['target']['positive']}” ÷ {entities} in the group",
            "signal groups = the lowest third against the highest third of a number column (a category: its highest-rate value against its lowest)",
        ] + (["risk score = a gradient-boosting model's out-of-fold probability; high ≥ 0.70, medium 0.30 – 0.70, low under 0.30",
              f"model tested on {m['test_n']:,} {entities} it never saw: average precision {m['ap']} (random {m['baseline_ap']})"] if m else [])
    else:
        title, subtitle = _title_for(None, entities, entry["source"], prof["measures"])
        spec["title"], spec["subtitle"] = name or title, subtitle
        spec["pages"] = _generic_pages(df, spec, prof, dims)
        spec["definitions"] = [f"{analysis.label_of(m)} = sum over the {entities} in view" for m in prof["measures"][:3]] + \
            ([f"periods follow {analysis.label_of(prof['date'])}; the latest month is compared with the month before and the same month a year earlier"] if prof["date"] else [])
    spec["definitions"] = [_cap(d) for d in spec.get("definitions") or []]
    spec["filters"] = _filters(df, spec, dims)
    spec["about"] = [f"{entry['source']} · {entry['table']} · {len(df):,} rows · {len(df.columns)} columns" + (" (a fair sample)" if sampled else ""),
                     ("Dates: " + analysis.label_of(prof["date"]) + ", so numbers compare periods") if prof.get("date")
                     else "No dates in the table, so comparisons are group against group, not period over period"]
    if spec.get("model", {}).get("leaks"):
        lk = spec["model"]["leaks"]
        names = ", ".join((c if len(c) <= 28 else c[:26] + "…") for c in lk[:2])
        spec["about"].append(f"{len(lk)} column{'s' if len(lk) != 1 else ''} left out of the model because {'they give' if len(lk) != 1 else 'it gives'} the answer away ({names})")
    return spec


def _filters(df: pd.DataFrame, spec: dict, dims: list[str]) -> list[dict]:
    out = []
    for c in dims[:4]:
        vals = [str(v) for v in df[c].dropna().astype(str).value_counts().index[:30]]
        out.append({"key": c, "label": analysis.label_of(c)[:1].upper() + analysis.label_of(c)[1:], "column": c, "values": sorted(vals, key=_natural)})
    if spec.get("model", {}).get("scores"):
        out.append({"key": "__risk", "label": "Risk tier", "column": None, "values": ["High", "Medium", "Low"]})
    if spec.get("dates"):
        out.append({"key": "__period", "label": "Period", "column": spec["dates"]["column"], "values": ["Last 12 months", "Last 3 months", "This year"]})
    return out


def _natural(v: str):
    """Ranges in their real order: "Less than $40K" < "$40K - $60K" < "$120K +"; words like Unknown last."""
    t = str(v).strip().lower()
    if t in ("unknown", "other", "none", "n/a", "na", ""):
        return (2, 0, t)
    m = re.match(r"[^\d-]*(-?\d[\d,]*\.?\d*)", t)
    if not m:
        return (1, 0, t)
    x = float(m.group(1).replace(",", ""))
    if re.match(r"\s*(less|under|below|<)", t):
        x -= 0.5
    return (0, x, t)


def _outcome_pages(df: pd.DataFrame, spec: dict, prof: dict, dims: list[str]) -> list[dict]:
    sigs = spec.get("signals") or []
    num_sigs = [s for s in sigs if s["kind"] == "numeric"]
    top = num_sigs[0]["column"] if num_sigs else None
    second = num_sigs[1]["column"] if len(num_sigs) > 1 else None
    tenure = next((c for c in prof["measures"] if TENURE.search(str(c))), None)
    has_scores = bool(spec.get("model", {}).get("scores"))
    ent, noun = spec["entities"], spec["outcome"]["noun"]
    ov = [{"id": "ov_rate", "kind": "o_rate", "span": 6, "h": 2, "hero": True, "curve": top}]
    ov.append({"id": "ov_count", "kind": "o_count", "span": 3})
    if top:
        ov.append({"id": "ov_m1", "kind": "o_measure", "span": 3, "column": top})
    if second:
        ov.append({"id": "ov_m2", "kind": "o_measure", "span": 3, "column": second})
    ov.append({"id": "ov_risk", "kind": "o_risk", "span": 3} if has_scores else {"id": "ov_m3", "kind": "o_measure", "span": 3,
                                                                                    "column": next((c for c in prof["measures"] if c not in (top, second)), top)})
    ov.append({"id": "ov_sig", "kind": "o_signals", "span": 6, "title": f"Where {noun} concentrates"})
    band_col = tenure if tenure and tenure not in (top, second) else (second or top)
    if spec.get("dates"):
        ov.append({"id": "ov_trend", "kind": "o_trend", "span": 6, "column": spec["dates"]["column"]})
    elif band_col:
        ov.append({"id": "ov_band", "kind": "o_band", "span": 6, "column": band_col})
    if has_scores:
        ov.append({"id": "ov_list", "kind": "o_list", "span": 7, "limit": 6})
    if dims:
        ov.append({"id": "ov_table", "kind": "o_table", "span": 5 if has_scores else 12, "column": dims[0]})
    pages = [{"id": "home", "name": "Home", "kind": "home"},
             {"id": "overview", "name": "Overview", "blurb": f"{noun[:1].upper() + noun[1:]}, usage and risk at a glance; who to call first", "blocks": ov}]
    sp = [{"id": "sg_imp", "kind": "o_importance", "span": 6}] if spec.get("model", {}).get("importance") else []
    sp.append({"id": "sg_sig", "kind": "o_signals", "span": 6 if sp else 12, "title": "Each signal against its comparison group"})
    for i, s in enumerate(num_sigs[:4]):
        sp.append({"id": f"sg_band{i}", "kind": "o_band", "span": 6, "column": s["column"]})
    pages.append({"id": "signals", "name": "Signals", "blurb": "The signals most associated with " + noun + ", and what the model weights", "blocks": sp})
    if dims:
        seg = []
        for i, d in enumerate(dims[:4]):
            seg.append({"id": f"seg_bars{i}", "kind": "o_dim", "span": 6, "column": d})
        seg.append({"id": "seg_table", "kind": "o_table", "span": 12, "column": dims[0]})
        pages.append({"id": "segments", "name": "Segments", "blurb": f"{noun[:1].upper() + noun[1:]} by " + ", ".join(analysis.label_of(d) for d in dims[:3]), "blocks": seg})
    if has_scores:
        pages.append({"id": "customers", "name": ent[:1].upper() + ent[1:], "blurb": f"Risk tiers and the {ent} to contact first",
                      "blocks": [{"id": "cu_tiers", "kind": "o_tiers", "span": 5}, {"id": "cu_list", "kind": "o_list", "span": 7, "limit": 25}]})
    return pages


def _generic_pages(df: pd.DataFrame, spec: dict, prof: dict, dims: list[str]) -> list[dict]:
    ms = prof["measures"][:4]
    date = spec.get("dates", {}).get("column")
    ov = []
    if ms:
        ov.append({"id": "ov_hero", "kind": "g_measure", "span": 6, "h": 2, "hero": True, "column": ms[0]})
    ov.append({"id": "ov_count", "kind": "g_count", "span": 3})
    small = [{"id": f"ov_m{i}", "kind": "g_measure", "span": 3, "column": m, "agg": "mean" if _unit(m) == "ratio" else "sum"} for i, m in enumerate(ms[1:3])]
    if ms:
        # the fourth tile: the main measure per day when there are dates, per row otherwise
        small.append({"id": "ov_rate", "kind": "g_measure", "span": 3, "column": ms[0], "agg": "per_day" if date else "mean"})
    ov += small[:3]
    if date and ms:
        ov.append({"id": "ov_trend", "kind": "g_trend", "span": 7, "column": ms[0]})
    if dims:
        ov.append({"id": "ov_bars", "kind": "g_dim", "span": 5 if date else 6, "column": dims[0], "measure": ms[0] if ms else None})
        ov.append({"id": "ov_table", "kind": "g_table", "span": 12, "column": dims[0]})
    pages = [{"id": "home", "name": "Home", "kind": "home"}, {"id": "overview", "name": "Overview", "blurb": "The main numbers and where they come from", "blocks": ov}]
    if len(dims) > 1:
        pages.append({"id": "breakdown", "name": "Breakdown", "blurb": "By " + ", ".join(analysis.label_of(d) for d in dims[:4]),
                      "blocks": [{"id": f"bd{i}", "kind": "g_dim", "span": 6, "column": d, "measure": ms[0] if ms else None} for i, d in enumerate(dims[:4])]})
    if date and ms:
        pages.append({"id": "trends", "name": "Trends", "blurb": "Month by month, against the year before",
                      "blocks": [{"id": f"tr{i}", "kind": "g_trend", "span": 6, "column": m} for i, m in enumerate(ms[:4])]})
    return pages


# ---- compute ------------------------------------------------------------------------------------

def _apply_filters(df: pd.DataFrame, spec: dict, filters: dict | None) -> tuple[pd.DataFrame, list[str]]:
    applied = []
    if not filters:
        return df, applied
    mask = pd.Series(True, index=df.index)
    for f in spec.get("filters") or []:
        v = filters.get(f["key"])
        if v in (None, "", "All"):
            continue
        if f["key"] == "__risk":
            scores = _scores(spec, df)
            if scores is None:
                continue
            lo, hi = {"High": (0.70, 1.01), "Medium": (0.30, 0.70), "Low": (-0.01, 0.30)}.get(v, (-1, 2))
            mask &= (scores >= lo) & (scores < hi)
        elif f["key"] == "__period":
            d = _dates(df, f["column"])
            end = d.max()
            if pd.isna(end):
                continue
            start = {"Last 12 months": end - pd.DateOffset(months=12), "Last 3 months": end - pd.DateOffset(months=3),
                     "This year": pd.Timestamp(year=end.year, month=1, day=1)}.get(v)
            if start is not None:
                mask &= d > start
        elif f.get("column") in df.columns:
            mask &= df[f["column"]].astype(str) == str(v)
        else:
            continue
        applied.append(f"{f['label']}: {v}")
    return df[mask], applied


def _scores(spec: dict, df: pd.DataFrame) -> pd.Series | None:
    m = spec.get("model") or {}
    sc = m.get("scores")
    if not sc or m.get("rows") != len(df):
        return None
    return pd.Series(sc, index=df.index, dtype=float)


def _y(df: pd.DataFrame, spec: dict) -> pd.Series:
    o = spec["outcome"]
    return (df[o["column"]].astype(str) == str(o["positive"])).astype(int)


def _rate(y: pd.Series) -> float | None:
    return _f(y.mean() * 100, 1) if len(y) else None


def _sig_mask(df: pd.DataFrame, s: dict, which: str = "group") -> pd.Series | None:
    c = s.get("column")
    if c not in df.columns:
        return None
    if s["kind"] == "categorical":
        v = s.get("value") if which == "group" else s.get("compare_value")
        return None if v is None else (df[c].astype(str) == str(v))
    th = s.get("threshold") or []
    if len(th) != 2:
        return None
    x = pd.to_numeric(df[c], errors="coerce")
    q1, q2 = th
    low = (x < q1) if q1 != q2 else (x <= q1)
    high = (x >= q2) if q1 != q2 else (x > q1)
    if s.get("direction") == "low":
        return low if which == "group" else high
    return high if which == "group" else low


def _bands(x: pd.Series, k: int = 6) -> tuple[pd.Series, list[str]] | None:
    xs = x.dropna()
    if xs.nunique() < 3:
        return None
    if xs.nunique() <= 8 and (xs % 1 == 0).all():
        vals = sorted(xs.unique())
        labels = [f"{int(v)}" for v in vals]
        return x.map({v: f"{int(v)}" for v in vals}), labels
    qs = sorted(set(np.quantile(xs, np.linspace(0, 1, k + 1)).tolist()))
    if len(qs) < 3:
        return None
    integer = bool((xs % 1 == 0).all())
    cats = pd.cut(x, bins=qs, include_lowest=True, duplicates="drop")
    labels = []
    for iv in cats.cat.categories:
        a, b = iv.left, iv.right
        labels.append(f"{_short(a, integer)}–{_short(b, integer)}")
    mapping = dict(zip(cats.cat.categories, labels))
    return cats.map(mapping).astype("object"), labels


def _short(v: float, integer: bool) -> str:
    v = 0.0 if abs(v) < 1e-9 else v
    if integer or float(v).is_integer():
        v = int(round(v))
        return f"{v / 1000:.1f}K".replace(".0K", "K") if abs(v) >= 10000 else f"{v:,}"
    return f"{v:.2f}" if abs(v) < 10 else f"{v:,.0f}"


def _block(df: pd.DataFrame, spec: dict, b: dict, base: pd.DataFrame) -> dict:
    kind = b["kind"]
    ent = spec["entities"]
    out = {"id": b["id"], "kind": kind, "span": b.get("span", 6), "h": b.get("h", 1), "hero": bool(b.get("hero"))}
    if b.get("title"):
        out["title"] = b["title"]
    if kind.startswith("o_"):
        o = spec["outcome"]
        noun, verb = o["noun"], o["verb"]
        y = _y(df, spec)
        all_rate = _rate(y)
        if kind == "o_rate":
            out.update(type="kpi", title=f"{noun[:1].upper() + noun[1:]} rate", sub=f"{ent[:1].upper() + ent[1:]} who {verb}, share of all",
                       value=all_rate, unit="pct", note=f"{int(y.sum()):,} of {len(y):,} {ent}", definition=spec["definitions"][0])
            chips = []
            sig = (spec.get("signals") or [None])[0]
            if sig:
                gm, cm = _sig_mask(df, sig), _sig_mask(df, sig, "compare")
                if gm is not None and gm.sum():
                    chips.append({"text": f"{_rate(y[gm])}% with {sig['group']}", "tone": "bad"})
                if cm is not None and cm.sum():
                    chips.append({"text": f"{_rate(y[cm])}% with {sig['compare']}", "tone": "good"})
            out["chips"] = chips
            col = b.get("curve")
            if col and col in df.columns:
                bd = _bands(pd.to_numeric(df[col], errors="coerce"), 5)
                if bd:
                    band, labels = bd
                    g = y.groupby(band).mean() * 100
                    out["curve"] = {"label": short(col), "points": [{"label": l, "value": _f(g.get(l), 1)} for l in labels if l in g.index]}
        elif kind == "o_count":
            out.update(type="kpi", title=ent[:1].upper() + ent[1:], sub="in view", value=int(len(df)), unit="count",
                       chips=[{"text": f"{int((y == 0).sum()):,} current", "tone": "neutral"}, {"text": f"{int(y.sum()):,} {verb}", "tone": "bad"}],
                       spark=_spark(df, spec, None))
            if not out.get("spark"):
                out["stack"] = [{"label": "Current", "value": int((y == 0).sum()), "tone": "accent"}, {"label": verb.capitalize(), "value": int(y.sum()), "tone": "bad"}]
        elif kind == "o_measure":
            c = b.get("column")
            x = pd.to_numeric(df[c], errors="coerce")
            a_, b_ = x[y == 1].mean(), x[y == 0].mean()
            worse = "bad"
            out.update(type="kpi", title=_cap(short(c)), sub=f"average per {ent[:-1] if ent.endswith('s') else ent}",
                       value=_f(x.mean(), 2), unit=_unit(c), definition=f"average of {c}",
                       chips=[{"text": f"{_fmt(a_, _unit(c))} among {verb}", "tone": worse}, {"text": f"{_fmt(b_, _unit(c))} among current", "tone": "good"}],
                       spark=_spark(df, spec, c))
        elif kind == "o_risk":
            sc = _scores(spec, base)
            if sc is None:
                out.update(type="empty", title="High risk now", text="Scores are available when the table has under 60,000 rows.")
            else:
                sc = sc.loc[df.index]
                cur = y == 0
                h = int(((sc >= 0.7) & cur).sum()); m = int(((sc >= 0.3) & (sc < 0.7) & cur).sum()); l = int(((sc < 0.3) & cur).sum())
                out.update(type="kpi", title="High risk now", sub="model score ≥ 0.70", value=h, unit="count", tone="bad",
                           chips=[{"text": f"{m:,} medium", "tone": "warn"}, {"text": f"{l:,} low", "tone": "neutral"}],
                           stack=[{"label": "High", "value": h, "tone": "bad"}, {"label": "Medium", "value": m, "tone": "warn"}, {"label": "Low", "value": l, "tone": "accent"}])
        elif kind == "o_signals":
            rows = []
            for s in (spec.get("signals") or [])[:6]:
                gm = _sig_mask(df, s)
                if gm is None or gm.sum() < 5:
                    continue
                rows.append({"label": s["group"], "value": _rate(y[gm]), "n": int(gm.sum()), "tone": "bad"})
            best = None
            for s in (spec.get("signals") or []):
                cm = _sig_mask(df, s, "compare")
                if cm is not None and cm.sum() >= 5:
                    r = _rate(y[cm])
                    if best is None or r < best["value"]:
                        best = {"label": s["compare"], "value": r, "n": int(cm.sum()), "tone": "good"}
            rows.sort(key=lambda r: -(r["value"] or 0))
            if best:
                rows.append(best)
            out.update(type="bars", title=b.get("title") or f"Where {noun} concentrates", sub=f"historical {noun} rate in each group · dashed line is all {ent} in view",
                       unit="pct", rows=rows, baseline={"value": all_rate, "label": f"All {all_rate}%"},
                       note="Groups overlap; rates are not additive.")
        elif kind == "o_band":
            c = b.get("column")
            bd = _bands(pd.to_numeric(df[c], errors="coerce"))
            if not bd:
                out.update(type="empty", title=analysis.label_of(c), text="Too few distinct values to band.")
            else:
                band, labels = bd
                g = y.groupby(band).agg(["mean", "size"])
                pts = [{"label": l, "value": _f(g.loc[l, "mean"] * 100, 1), "n": int(g.loc[l, "size"])} for l in labels if l in g.index and g.loc[l, "size"] >= 5]
                lab = short(c)
                out.update(type="line", title=f"{noun[:1].upper() + noun[1:]} rate by {lab}", sub=f"{lab}, in bands of about equal size · % who {verb}",
                           unit="pct", points=pts, baseline={"value": all_rate, "label": f"All {all_rate}%"})
        elif kind == "o_trend":
            d = _dates(df, b["column"])
            per = d.dt.to_period("M")
            g = y.groupby(per).agg(["mean", "size"]).sort_index()
            g = g[g["size"] >= 10].tail(18)
            pts = [{"label": p.strftime("%b %y"), "value": _f(r["mean"] * 100, 1), "n": int(r["size"])} for p, r in g.iterrows()]
            out.update(type="line", title=f"{noun[:1].upper() + noun[1:]} rate by month", sub=f"by {analysis.label_of(b['column'])} · % who {verb}",
                       unit="pct", points=pts, baseline={"value": all_rate, "label": f"All {all_rate}%"})
        elif kind == "o_dim":
            c = b["column"]
            s = df[c].astype(str).where(df[c].notna())
            g = y.groupby(s).agg(["mean", "size", "sum"])
            g = g[g["size"] >= 5].sort_values("size", ascending=False).head(10)
            margin = max(1.0, 0.1 * (all_rate or 0))

            def tone_of(rate: float, n: int) -> str:
                if n < 30:
                    return "neutral"
                return "bad" if rate > (all_rate or 0) + margin else "good" if rate < (all_rate or 0) - margin else "neutral"
            rows = [{"label": str(k), "value": _f(r["mean"] * 100, 1), "n": int(r["size"]), "count": int(r["sum"]), "small": int(r["size"]) < 30,
                     "tone": tone_of(r["mean"] * 100, int(r["size"]))} for k, r in g.iterrows()]
            rows.sort(key=lambda r: _natural(r["label"]))
            lab = analysis.label_of(c)
            out.update(type="bars", title=f"{noun[:1].upper() + noun[1:]} by {lab}",
                       sub=f"% who {verb} · coloured when more than {margin:.1f} points from all {ent}",
                       unit="pct", rows=rows, baseline={"value": all_rate, "label": f"All {all_rate}%"}, show_n=True, n_label=ent)
        elif kind == "o_table":
            c = b["column"]
            s = df[c].astype(str).where(df[c].notna())
            top_m = [x["column"] for x in (spec.get("signals") or []) if x["kind"] == "numeric"][:2]
            g = pd.DataFrame({"n": y.groupby(s).size(), "pos": y.groupby(s).sum()})
            for m in top_m:
                g[m] = pd.to_numeric(df[m], errors="coerce").groupby(s).mean()
            g = g[g["n"] >= 1].sort_values("n", ascending=False).head(12)
            cols = [{"key": "label", "label": analysis.label_of(c)[:1].upper() + analysis.label_of(c)[1:], "align": "left"},
                    {"key": "n", "label": ent[:1].upper() + ent[1:], "unit": "count"}, {"key": "pos", "label": verb[:1].upper() + verb[1:], "unit": "count"},
                    {"key": "rate", "label": "Rate", "unit": "pct"}]
            cols += [{"key": m, "label": "Avg " + short(m), "unit": _unit(m)} for m in top_m]
            cols.append({"key": "vs", "label": "vs all", "unit": "pp"})
            rows = []
            for k, r in g.iterrows():
                rate = _f(r["pos"] / r["n"] * 100, 1) if r["n"] else None
                row = {"label": str(k), "n": int(r["n"]), "pos": int(r["pos"]), "rate": rate, "vs": _f((rate or 0) - (all_rate or 0), 1)}
                for m in top_m:
                    row[m] = _f(r[m], 2)
                rows.append(row)
            total = {"label": "All", "n": int(len(y)), "pos": int(y.sum()), "rate": all_rate, "vs": None}
            for m in top_m:
                total[m] = _f(pd.to_numeric(df[m], errors="coerce").mean(), 2)
            out.update(type="table", title=f"By {analysis.label_of(c)}", sub=f"vs all = rate minus {all_rate}%", columns=cols, rows=rows, total=total)
        elif kind == "o_list":
            sc = _scores(spec, base)
            if sc is None:
                out.update(type="empty", title=f"{ent[:1].upper() + ent[1:]} to contact first", text="Scores are available when the table has under 60,000 rows.")
            else:
                sc = sc.loc[df.index]
                cur = sc[y == 0].sort_values(ascending=False)
                lim = int(b.get("limit") or 10)
                idc = (spec.get("model") or {}).get("id_column")
                sigs = (spec.get("signals") or [])[:4]
                rows = []
                for idx in cur.index[:lim]:
                    tags = []
                    for s in sigs:
                        gm = _sig_mask(df.loc[[idx]], s)
                        if gm is not None and bool(gm.iloc[0]):
                            v = df.at[idx, s["column"]]
                            tags.append(f"{short(s['column'])} {_fmt(v, _unit(s['column']))}" if s["kind"] == "numeric" else str(v))
                    meta = " · ".join(str(df.at[idx, d]) for d in [f["column"] for f in spec.get("filters") or [] if f.get("column") and not f["key"].startswith("__")][:2])
                    rows.append({"id": str(df.at[idx, idc]) if idc else str(idx), "score": _f(sc[idx], 2), "meta": meta, "tags": tags[:3]})
                total_high = int((cur >= 0.7).sum())
                out.update(type="list", title=b.get("title") or f"{ent[:1].upper() + ent[1:]} to contact first", sub=f"high-risk current {ent} by score · {total_high:,} in the high tier",
                           rows=rows, total=total_high, id_label=_id_label(idc))
        elif kind == "o_importance":
            imp = (spec.get("model") or {}).get("importance") or []
            out.update(type="bars", title="What the model weights", sub="held-out precision lost when the signal is shuffled · not causal", unit="num",
                       rows=[{"label": d["label"], "value": d["value"], "tone": "accent"} for d in imp[:8]], decimals=3)
        elif kind == "o_tiers":
            sc = _scores(spec, base)
            if sc is None:
                out.update(type="empty", title="Risk tiers", text="Scores are available when the table has under 60,000 rows.")
            else:
                sc = sc.loc[df.index]
                tiers = []
                for key, lab, lo, hi in (("high", "High risk", 0.7, 1.01), ("medium", "Medium risk", 0.3, 0.7), ("low", "Low risk", -0.01, 0.3)):
                    band = (sc >= lo) & (sc < hi)
                    tiers.append({"key": key, "label": lab, "count": int((band & (y == 0)).sum()), "hist": _rate(y[band]) if band.sum() else None,
                                  "rule": "score ≥ 0.70" if key == "high" else "score 0.30 – 0.70" if key == "medium" else "score under 0.30"})
                out.update(type="tiers", title="Current " + ent + " by risk tier", sub=f"historical {noun} rate of each band shown under it", tiers=tiers)
        return out
    # generic
    if kind == "g_count":
        out.update(type="kpi", title=ent[:1].upper() + ent[1:], sub="in view", value=int(len(df)), unit="count", spark=_spark(df, spec, None, agg="count"),
                   chips=_period_chips(df, spec, None, "count"))
    elif kind == "g_measure":
        c = b["column"]
        agg = b.get("agg") or "sum"
        x = pd.to_numeric(df[c], errors="coerce")
        title = analysis.label_of(c)[:1].upper() + analysis.label_of(c)[1:]
        if agg == "per_day" and spec.get("dates"):
            d = _dates(df, spec["dates"]["column"]).dt.normalize()
            daily = x.groupby(d).sum()
            out.update(type="kpi", title=f"{title} per day", sub="average day", value=_f(daily.mean(), 2), unit=_unit(c),
                       spark=[_f(v, 2) for v in daily.tail(30).tolist()] if len(daily) > 3 else None,
                       chips=[{"text": f"best {_fmt(daily.max(), _unit(c))}", "tone": "good"}, {"text": f"{len(daily):,} days", "tone": "neutral"}],
                       definition=f"sum of {c} per day, averaged over {len(daily):,} days")
        else:
            val = x.sum() if agg == "sum" else x.mean()
            out.update(type="kpi", title=title if agg == "sum" else f"{title} per {ent[:-1] if ent.endswith('s') else ent}", sub="total" if agg == "sum" else "average",
                       value=_f(val, 2), unit=_unit(c), spark=_spark(df, spec, c, agg=agg), chips=_period_chips(df, spec, c, agg),
                       definition=f"{agg} of {c}")
        if b.get("hero") and spec.get("dates"):
            pts = _monthly(df, spec, c, agg)
            out["curve"] = {"label": "by month", "points": pts[-12:]}
    elif kind == "g_trend":
        c = b["column"]
        pts = _monthly(df, spec, c, "sum")
        prev = _monthly(df, spec, c, "sum", shift_year=True)
        lab = analysis.label_of(c)
        out.update(type="line", title=f"{lab[:1].upper() + lab[1:]} by month", sub="last 12 months · dashed line is the year before", unit=_unit(c),
                   points=pts[-12:], compare=prev[-12:] if prev else None)
    elif kind == "g_dim":
        c, m = b["column"], b.get("measure")
        s = df[c].astype(str).where(df[c].notna())
        if m:
            g = pd.to_numeric(df[m], errors="coerce").groupby(s).sum().sort_values(ascending=False).head(10)
            tot = float(g.sum()) or 1.0
            rows = [{"label": str(k), "value": _f(v, 2), "share": _f(v / tot * 100, 1), "tone": "accent"} for k, v in g.items()]
            lab = analysis.label_of(m)
            out.update(type="bars", title=f"{lab[:1].upper() + lab[1:]} by {analysis.label_of(c)}", sub="total · label shows share of all", unit=_unit(m), rows=rows, show_share=True)
        else:
            g = s.value_counts().head(10)
            rows = [{"label": str(k), "value": int(v), "tone": "accent"} for k, v in g.items()]
            out.update(type="bars", title=f"{ent[:1].upper() + ent[1:]} by {analysis.label_of(c)}", sub="count", unit="count", rows=rows)
    elif kind == "g_table":
        c = b["column"]
        s = df[c].astype(str).where(df[c].notna())
        ms = [m for m in (spec.get("_measures") or [])] or [x["column"] for p in spec["pages"] for x in p.get("blocks", []) if x["kind"] == "g_measure"][:3]
        g = pd.DataFrame({"n": s.value_counts()})
        for m in ms:
            g[m] = pd.to_numeric(df[m], errors="coerce").groupby(s).sum()
        g = g.sort_values("n", ascending=False).head(12)
        cols = [{"key": "label", "label": analysis.label_of(c)[:1].upper() + analysis.label_of(c)[1:], "align": "left"}, {"key": "n", "label": "Rows", "unit": "count"}]
        cols += [{"key": m, "label": analysis.label_of(m)[:1].upper() + analysis.label_of(m)[1:], "unit": _unit(m)} for m in ms]
        main = ms[0] if ms else None
        if main:
            cols.append({"key": "share", "label": "Share", "unit": "pct"})
        tot_main = float(pd.to_numeric(df[main], errors="coerce").sum()) if main else 0
        rows = []
        for k, r in g.iterrows():
            row = {"label": str(k), "n": int(r["n"])}
            for m in ms:
                row[m] = _f(r[m], 2)
            if main:
                row["share"] = _f(r[main] / tot_main * 100, 1) if tot_main else None
            rows.append(row)
        total = {"label": "All", "n": int(len(df))}
        for m in ms:
            total[m] = _f(pd.to_numeric(df[m], errors="coerce").sum(), 2)
        if main:
            total["share"] = 100.0
        out.update(type="table", title=f"By {analysis.label_of(c)}", sub="totals in view", columns=cols, rows=rows, total=total)
    return out


def _id_label(col: str | None) -> str:
    """CLIENTNUM -> "Client", customer_id -> "Customer", booking_ref -> "Booking"; "Row" without an id column."""
    if not col:
        return "Row"
    base = re.sub(r"(num(ber)?|_?id|_?no|_?ref|_?key|_?code)$", "", str(col), flags=re.I).strip("_ ") or str(col)
    return _cap(analysis.label_of(base))


def _fmt(v, unit: str) -> str:
    try:
        x = float(v)
    except (TypeError, ValueError):
        return str(v)
    if not math.isfinite(x):
        return "—"
    if unit == "money":
        return f"${x:,.0f}" if abs(x) >= 100 else f"${x:,.2f}"
    if unit == "ratio":
        return f"{x:.2f}"
    if unit == "pct":
        return f"{x:.1f}%"
    return f"{x:,.0f}" if abs(x) >= 100 or float(x).is_integer() else f"{x:,.1f}"


def _spark(df: pd.DataFrame, spec: dict, col: str | None, agg: str = "mean") -> list | None:
    """A sparkline: by month when there are dates; else the measure's distribution (sorted deciles)."""
    if spec.get("dates"):
        pts = _monthly(df, spec, col, agg if col else "count")
        return [p["value"] for p in pts[-12:]] if len(pts) >= 3 else None
    if col is None:
        return None
    x = pd.to_numeric(df[col], errors="coerce").dropna()
    if len(x) < 20:
        return None
    return [_f(v, 3) for v in np.quantile(x, np.linspace(0.05, 0.95, 10))]


def _monthly(df: pd.DataFrame, spec: dict, col: str | None, agg: str, shift_year: bool = False) -> list[dict]:
    """Totals by calendar month; a last month the data has not finished (it stops 3+ days before the month ends)
    is left out, so "this month vs last month" never compares a part month with a whole one."""
    d = _dates(df, spec["dates"]["column"])
    end = d.max()
    if not pd.isna(end) and (end + pd.Timedelta(days=3)).month == end.month:
        d = d.where(d.dt.to_period("M") != end.to_period("M"))
    per = d.dt.to_period("M")
    if col is None or agg == "count":
        g = per.value_counts().sort_index()
    else:
        x = pd.to_numeric(df[col], errors="coerce")
        g = x.groupby(per).sum() if agg == "sum" else x.groupby(per).mean()
        g = g.sort_index()
    if g.empty:
        return []
    g = g[g.index.notna()]
    if shift_year:
        last = g.index.max()
        months = pd.period_range(last - 11, last, freq="M")
        return [{"label": p.strftime("%b %y"), "value": _f(g.get(p - 12), 2)} for p in months] if any((p - 12) in g.index for p in months) else []
    return [{"label": p.strftime("%b %y"), "value": _f(v, 2)} for p, v in g.items()]


COST = re.compile(r"discount|cost|refund|return|churn|cancel|complaint|defect|delay|tax|fee|expense|cogs", re.I)


def _period_chips(df: pd.DataFrame, spec: dict, col: str | None, agg: str) -> list[dict]:
    if not spec.get("dates"):
        return []
    pts = _monthly(df, spec, col, agg)
    if len(pts) < 2:
        return []
    by = {p["label"]: p["value"] for p in pts}
    cur = pts[-1]
    prev = pts[-2]
    chips = []
    try:
        last = pd.Period(datetime.strptime(cur["label"], "%b %y"), freq="M")
        py = by.get((last - 12).strftime("%b %y"))
    except ValueError:
        py = None
    for label, base in (("PM", prev["value"]), ("PY", py)):
        if base and cur["value"] is not None:
            ch = (cur["value"] - base) / abs(base) * 100
            up_is_good = not (col and COST.search(str(col)))
            chips.append({"text": f"{'▲' if ch >= 0 else '▼'} {abs(ch):.1f}% {label}", "tone": ("good" if ch >= 0 else "bad") if up_is_good else ("bad" if ch >= 0 else "good")})
    if chips:
        chips.insert(0, {"text": cur["label"], "tone": "neutral"})
    return chips


def compute(db: Session, user: models.User, dash: models.Dashboard, page_id: str | None, filters: dict | None) -> dict:
    spec = dash.board_spec or {}
    src = spec.get("source") or {}
    base, sampled = frame(db, user, src["source_id"], src.get("table"))
    page = next((p for p in spec.get("pages") or [] if p["id"] == page_id), None) or (spec.get("pages") or [{}])[0]
    df, applied = _apply_filters(base, spec, filters or spec.get("default_filters"))
    out = {"id": dash.id, "name": dash.name, "page": {"id": page.get("id"), "name": page.get("name")}, "rows_in_view": int(len(df)),
           "rows_total": int(len(base)), "applied": applied, "refreshed_at": datetime.utcnow().isoformat() + "Z", "blocks": []}
    if page.get("kind") == "home":
        out["home"] = _home(df, spec, base)
        return _clean(out)
    if len(df) == 0:
        out["empty"] = "No rows match these filters."
        return _clean(out)
    for b in page.get("blocks") or []:
        try:
            out["blocks"].append(_block(df, spec, b, base))
        except Exception as e:  # noqa: BLE001 - one block never breaks the page
            print(f"[board] block {b.get('id')} failed: {e}")
            out["blocks"].append({"id": b["id"], "type": "empty", "span": b.get("span", 6), "title": b.get("title") or "This block",
                                  "text": "This block could not be computed on the current data."})
    return _clean(out)


def _home(df: pd.DataFrame, spec: dict, base: pd.DataFrame) -> dict:
    kpis = []
    ent = spec["entities"]
    if spec.get("outcome"):
        y = _y(df, spec)
        o = spec["outcome"]
        kpis.append({"label": ent[:1].upper() + ent[1:], "value": int(len(df)), "unit": "count", "note": f"{int(y.sum()):,} {o['verb']}"})
        kpis.append({"label": f"{o['noun'][:1].upper() + o['noun'][1:]} rate", "value": _rate(y), "unit": "pct", "note": "of all in view"})
        sig = next((s for s in spec.get("signals") or [] if s["kind"] == "numeric"), None)
        if sig:
            x = pd.to_numeric(df[sig["column"]], errors="coerce")
            kpis.append({"label": _cap(short(sig["column"])), "value": _f(x.mean(), 2), "unit": _unit(sig["column"]),
                         "note": f"{_fmt(x[y == 1].mean(), _unit(sig['column']))} among {o['verb']}"})
        sc = _scores(spec, base)
        if sc is not None:
            sc = sc.loc[df.index]
            kpis.append({"label": "High risk now", "value": int(((sc >= 0.7) & (y == 0)).sum()), "unit": "count",
                         "note": f"{int(((sc >= 0.3) & (sc < 0.7) & (y == 0)).sum()):,} medium"})
    else:
        kpis.append({"label": ent[:1].upper() + ent[1:], "value": int(len(df)), "unit": "count", "note": "rows in view"})
        for p in spec.get("pages") or []:
            for b in p.get("blocks") or []:
                if b["kind"] == "g_measure" and b.get("agg") != "per_day" and len(kpis) < 4:
                    x = pd.to_numeric(df[b["column"]], errors="coerce")
                    agg = b.get("agg") or "sum"
                    kpis.append({"label": analysis.label_of(b["column"]).capitalize(), "value": _f(x.sum() if agg == "sum" else x.mean(), 2),
                                 "unit": _unit(b["column"]), "note": "total" if agg == "sum" else "average"})
            if len(kpis) >= 4:
                break
    pages = [{"id": p["id"], "name": p["name"], "blurb": p.get("blurb") or "", "kinds": [b["kind"] for b in p.get("blocks") or []][:6]}
             for p in spec.get("pages") or [] if p.get("kind") != "home"]
    return {"title": spec.get("title"), "subtitle": spec.get("subtitle"), "kpis": kpis, "pages": pages, "definitions": spec.get("definitions") or [],
            "about": spec.get("about") or [], "theme": spec.get("theme")}


# ---- edits by talking ----------------------------------------------------------------------------

EDIT = re.compile(
    r"\b(make|turn|switch|set|change|use|go)\b.{0,30}\b(dark|light|navy|forest|graphite|sand|ink|theme|brand colou?rs?|colou?rs?|accent)\b"
    r"|#[0-9a-f]{6}\b|\b(dark|light) (mode|theme)\b"
    r"|\b(rename|call|name|title)\s+(it|this|the dashboard|dashboard)\b"
    r"|\badd\b.{0,50}\b(chart|graph|kpi|table|block|tile|breakdown|view)\b|\badd\s+\w+(\s+\w+)?\s+by\s+\w+"
    r"|\b(remove|delete|drop|hide)\b.{0,50}\b(chart|graph|kpi|table|block|tile|card|list|breakdown)\b|\b(remove|delete|drop|hide) the\b"
    r"|\b(filter|limit)\s+(it|this|the dashboard|dashboard)?\s*(to|by)\b|\b(show only|only show)\b|^\s*undo\b", re.I)


def wants_edit(text: str) -> bool:
    return bool(EDIT.search(text or ""))


def apply_edit(db: Session, user: models.User, dash: models.Dashboard, text: str, columns: list[str] | None = None,
               numeric: set[str] | None = None) -> dict:
    """Returns {changes[], page_id?}. Raises ValueError with a plain sentence when nothing matched."""
    spec = copy.deepcopy(dash.board_spec or {})
    t = (text or "").strip()
    tl = t.lower()
    changes: list[str] = []
    page_id = None
    history = list(spec.get("_history") or [])
    if re.match(r"^\s*undo\b", tl):
        if not history:
            raise ValueError("There is nothing to undo on this dashboard yet.")
        prev = history.pop()
        restored = {k: v for k, v in prev.items() if k != "_label"}
        restored["model"] = spec.get("model")
        restored["_history"] = history
        dash.board_spec = restored
        dash.name = (restored.get("title") or dash.name)[:80]
        flag_modified(dash, "board_spec")
        db.commit()
        return {"changes": [f"undid “{prev.get('_label') or 'the last change'}”"], "page_id": None}
    snapshot = copy.deepcopy({k: v for k, v in spec.items() if k not in ("model", "_history")})
    m = re.search(r"\b(?:rename(?: it)?(?: to)?|call it|name it|title it)\s+[\"“']?(.+?)[\"”']?\s*$", t, re.I)
    if m:
        new = m.group(1).strip().rstrip(".")
        new = re.split(r"\s+and\s+(?:make|use|switch|change)\b", new, flags=re.I)[0].strip()
        if new:
            dash.name = new[:80]
            spec["title"] = new[:80]
            changes.append(f"renamed to “{new[:80]}”")
    hexm = re.search(r"#([0-9a-f]{6})\b", tl)
    if hexm and re.search(r"\b(brand|colou?r|accent)\b", tl):
        spec["accent"] = "#" + hexm.group(1)
        changes.append(f"accent colour set to #{hexm.group(1).upper()}")
    else:
        for key, pat in THEME_WORDS.items():
            if re.search(pat, tl) and re.search(r"\b(make|switch|change|use|set|theme|turn|go)\b", tl):
                if key == "graphite" and re.search(r"\b(navy|forest|blue|green)\b", tl):
                    continue
                spec["theme"] = key
                spec.pop("accent", None)
                changes.append(f"theme set to {THEMES[key]['label']}")
                break
    rm = re.search(r"\b(?:remove|delete|drop|hide)\s+(?:the\s+)?(.+?)(?:\s+(?:block|chart|tile|card|table))?\s*$", t, re.I)
    if rm and not changes:
        want = rm.group(1).lower()
        for p in spec.get("pages") or []:
            for b in list(p.get("blocks") or []):
                lab = " ".join(filter(None, [b.get("title"), analysis.label_of(b.get("column") or ""), b["kind"].replace("o_", "").replace("g_", "")])).lower()
                if want and (want in lab or all(w in lab for w in re.findall(r"[a-z]{4,}", want))):
                    p["blocks"].remove(b)
                    changes.append(f"removed “{b.get('title') or analysis.label_of(b.get('column') or b['kind'])}” from {p['name']}")
                    page_id = p["id"]
                    break
            if changes:
                break
    add = re.search(r"\badd\b\s+(?:a\s+|an\s+)?(?:chart|graph|block|breakdown|view|table)?\s*(?:of|for|showing|with)?\s*(.+)$", t, re.I)
    if add and not changes and columns:
        want = add.group(1)
        numeric = numeric or set()
        dims = [c for c in columns if c not in numeric]
        left, _, right = want.partition(" by ")
        dim_text = right or want

        def best(text_: str, cands: list[str]) -> str | None:
            sc = sorted(((analysis.word_score(text_, c), i, c) for i, c in enumerate(cands)), key=lambda x: (-x[0], x[1]))
            return sc[0][2] if sc and sc[0][0] > 0 else None
        col = best(dim_text, dims) or best(dim_text, columns)
        measure = best(left, [c for c in columns if c in numeric]) if right else None
        if col:
            target = next((p for p in spec.get("pages") or [] if p["id"] in ("segments", "breakdown")), None) or \
                next((p for p in spec.get("pages") or [] if p["id"] == "overview"), None)
            if target is not None:
                kind = ("o_dim" if spec.get("outcome") else "g_dim")
                bid = f"add{int(time.time() * 1000) % 10_000_000}"
                blk = {"id": bid, "kind": kind, "span": 6, "column": col}
                if not spec.get("outcome"):
                    blk["measure"] = measure or next((b.get("column") for p in spec["pages"] for b in p.get("blocks") or [] if b["kind"] == "g_measure"), None)
                target.setdefault("blocks", []).insert(0, blk)
                what = f"{analysis.label_of(blk['measure'])} by {analysis.label_of(col)}" if blk.get("measure") else analysis.label_of(col)
                changes.append(f"added {what} to {target['name']}")
                page_id = target["id"]
    fm = re.search(r"\b(?:filter|show only|only show|limit it)\s+(?:to|by)?\s*(.+)$", t, re.I)
    if fm and not changes:
        want = fm.group(1).strip().lower().rstrip(".")
        for f in spec.get("filters") or []:
            v = next((v for v in f["values"] if v.lower() == want or v.lower() in want), None)
            if v:
                spec["default_filters"] = {**(spec.get("default_filters") or {}), f["key"]: v}
                changes.append(f"filtered to {f['label']}: {v}")
                break
    if not changes:
        raise ValueError("I can change the theme (Ink, Forest, Navy, Graphite, Sand, or a brand colour like #0A7C5A), rename it, "
                         "add a breakdown (\"add churn by education level\"), remove a block, or set a default filter (\"filter to Blue\").")
    snapshot["_label"] = "; ".join(changes)
    spec["_history"] = (history + [snapshot])[-10:]
    dash.board_spec = spec
    flag_modified(dash, "board_spec")
    db.commit()
    return {"changes": changes, "page_id": page_id}


# ---- create ---------------------------------------------------------------------------------------

def create(db: Session, user: models.User, conv: models.Conversation | None, source_id: str, table: str | None, *, question: str = "",
           analysis_result: dict | None = None, name: str | None = None, theme: str = "ink") -> models.Dashboard:
    spec = design(db, user, source_id, table, question=question, analysis_result=analysis_result, theme=theme, name=name)
    ds = db.get(models.DataSource, source_id)
    dash = models.Dashboard(owner_id=user.id, workspace_id=(ds.workspace_id if ds else None) or (conv.workspace_id if conv else None),
                            name=spec["title"][:80], layout_version=4, source_conversation_id=conv.id if conv else None)
    if hasattr(models.Dashboard, "datasource_id"):
        dash.datasource_id = source_id
    dash.board_spec = spec
    db.add(dash)
    db.commit()
    db.refresh(dash)
    return dash


def summary(dash: models.Dashboard) -> dict:
    """The spec as the page needs it (no per-row scores)."""
    spec = dict(dash.board_spec or {})
    model = dict(spec.get("model") or {})
    model.pop("scores", None)
    th = THEMES.get(spec.get("theme") or "ink", THEMES["ink"]) | ({"accent": spec["accent"]} if spec.get("accent") else {})
    return _clean({"id": dash.id, "name": dash.name, "title": spec.get("title") or dash.name, "subtitle": spec.get("subtitle"),
                   "theme": spec.get("theme") or "ink", "palette": th, "themes": {k: v["label"] for k, v in THEMES.items()},
                   "source": spec.get("source"), "entities": spec.get("entities"), "outcome": spec.get("outcome"), "dates": spec.get("dates"),
                   "filters": spec.get("filters") or [], "default_filters": spec.get("default_filters") or {},
                   "pages": [{"id": p["id"], "name": p["name"], "blurb": p.get("blurb"), "kind": p.get("kind") or "page"} for p in spec.get("pages") or []],
                   "definitions": spec.get("definitions") or [], "about": spec.get("about") or [], "model": model or None,
                   "built_at": spec.get("built_at"), "sampled": spec.get("sampled"), "conversation_id": dash.source_conversation_id})


# ---- exports from the same computed pages ---------------------------------------------------------

def _rgb(hexs: str) -> tuple[int, int, int]:
    h = hexs.lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def _num_fmt(unit: str | None) -> str:
    return {"money": '"$"#,##0', "pct": '0.0"%"', "pp": '+0.0" pp";-0.0" pp"', "ratio": "0.00", "count": "#,##0"}.get(unit or "", "#,##0.##")


def workbook(db: Session, user: models.User, dash: models.Dashboard, filters: dict | None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    spec = dash.board_spec or {}
    th = THEMES.get(spec.get("theme") or "ink", THEMES["ink"])
    accent = (spec.get("accent") or th["accent"]).lstrip("#")
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "Home"
    home = compute(db, user, dash, "home", filters).get("home") or {}
    ws0["A1"] = spec.get("title") or dash.name
    ws0["A1"].font = Font(size=20, bold=True)
    ws0["A2"] = spec.get("subtitle") or ""
    ws0["A2"].font = Font(size=11, color="5B6472")
    r = 4
    for k in home.get("kpis") or []:
        ws0.cell(row=r, column=1, value=k["label"]).font = Font(bold=True)
        c = ws0.cell(row=r, column=2, value=k["value"])
        c.number_format = _num_fmt(k.get("unit"))
        ws0.cell(row=r, column=3, value=k.get("note") or "").font = Font(color="5B6472")
        r += 1
    r += 1
    ws0.cell(row=r, column=1, value="Definitions").font = Font(bold=True)
    for d in home.get("definitions") or []:
        r += 1
        ws0.cell(row=r, column=1, value=d)
    r += 2
    ws0.cell(row=r, column=1, value="About the data").font = Font(bold=True)
    for a in home.get("about") or []:
        r += 1
        ws0.cell(row=r, column=1, value=a)
    ws0.column_dimensions["A"].width = 46
    ws0.column_dimensions["B"].width = 16
    ws0.column_dimensions["C"].width = 30
    head_fill = PatternFill("solid", fgColor=accent)
    for p in spec.get("pages") or []:
        if p.get("kind") == "home":
            continue
        out = compute(db, user, dash, p["id"], filters)
        ws = wb.create_sheet(p["name"][:30])
        ws["A1"] = f"{spec.get('title') or dash.name} · {p['name']}"
        ws["A1"].font = Font(size=16, bold=True)
        ws["A2"] = (f"{out.get('rows_in_view', 0):,} rows in view" + (" · " + ", ".join(out.get("applied") or []) if out.get("applied") else ""))
        ws["A2"].font = Font(color="5B6472")
        row = 4
        for b in out.get("blocks") or []:
            t = b.get("type")
            ws.cell(row=row, column=1, value=b.get("title") or "").font = Font(size=12, bold=True)
            if b.get("sub"):
                ws.cell(row=row, column=2, value=b.get("sub")).font = Font(color="8A93A3", italic=True)
            row += 1
            if t == "kpi":
                c = ws.cell(row=row, column=1, value=b.get("value"))
                c.number_format = _num_fmt(b.get("unit"))
                c.font = Font(size=14, bold=True)
                ws.cell(row=row, column=2, value=" · ".join(ch["text"] for ch in b.get("chips") or []))
                row += 2
            elif t in ("bars", "line"):
                items = b.get("rows") or b.get("points") or []
                ws.cell(row=row, column=1, value="Group" if t == "bars" else "Band").font = Font(bold=True, color="FFFFFF")
                ws.cell(row=row, column=2, value="Value").font = Font(bold=True, color="FFFFFF")
                ws.cell(row=row, column=3, value="Rows").font = Font(bold=True, color="FFFFFF")
                for col in (1, 2, 3):
                    ws.cell(row=row, column=col).fill = head_fill
                row += 1
                for it in items:
                    ws.cell(row=row, column=1, value=it.get("label"))
                    c = ws.cell(row=row, column=2, value=it.get("value"))
                    c.number_format = _num_fmt(b.get("unit"))
                    ws.cell(row=row, column=3, value=it.get("n"))
                    row += 1
                if b.get("baseline"):
                    ws.cell(row=row, column=1, value=b["baseline"].get("label")).font = Font(italic=True, color="5B6472")
                    row += 1
                row += 1
            elif t == "table":
                cols = b.get("columns") or []
                for j, ccol in enumerate(cols, start=1):
                    cell = ws.cell(row=row, column=j, value=ccol["label"])
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = head_fill
                row += 1
                first = row
                for rr in b.get("rows") or []:
                    for j, ccol in enumerate(cols, start=1):
                        cell = ws.cell(row=row, column=j, value=rr.get(ccol["key"]))
                        if ccol.get("unit"):
                            cell.number_format = _num_fmt(ccol.get("unit"))
                    row += 1
                last = row - 1
                tot = b.get("total") or {}
                for j, ccol in enumerate(cols, start=1):
                    letter = get_column_letter(j)
                    if j == 1:
                        v = "Total"
                    elif ccol.get("unit") == "count" and last >= first:
                        v = f"=SUM({letter}{first}:{letter}{last})"
                    else:
                        v = tot.get(ccol["key"])
                    cell = ws.cell(row=row, column=j, value=v)
                    cell.font = Font(bold=True)
                    if ccol.get("unit"):
                        cell.number_format = _num_fmt(ccol.get("unit"))
                row += 2
            elif t == "list":
                for j, h in enumerate([b.get("id_label") or "Row", "Score", "Signals", "Segment"], start=1):
                    cell = ws.cell(row=row, column=j, value=h)
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = head_fill
                row += 1
                for it in b.get("rows") or []:
                    ws.cell(row=row, column=1, value=it.get("id"))
                    ws.cell(row=row, column=2, value=it.get("score")).number_format = "0.00"
                    ws.cell(row=row, column=3, value=" · ".join(it.get("tags") or []))
                    ws.cell(row=row, column=4, value=it.get("meta"))
                    row += 1
                row += 1
            elif t == "tiers":
                for it in b.get("tiers") or []:
                    ws.cell(row=row, column=1, value=it["label"])
                    ws.cell(row=row, column=2, value=it["count"]).number_format = "#,##0"
                    ws.cell(row=row, column=3, value=f"{it.get('hist')}% historically · {it['rule']}")
                    row += 1
                row += 1
            else:
                ws.cell(row=row, column=1, value=b.get("text") or "")
                row += 2
        ws.column_dimensions["A"].width = 40
        for col in "BCDEFGH":
            ws.column_dimensions[col].width = 16
        for rr in ws.iter_rows(min_row=1, max_row=3):
            for c in rr:
                c.alignment = Alignment(vertical="center")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def deck(db: Session, user: models.User, dash: models.Dashboard, filters: dict | None) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Emu, Inches, Pt
    spec = dash.board_spec or {}
    th = THEMES.get(spec.get("theme") or "ink", THEMES["ink"])
    accent = spec.get("accent") or th["accent"]
    bg, card, ink, warn = th["bg"], th["card"], th["ink"], th["warn"]
    muted = "#8A93A3" if th["mode"] == "light" else "#A9B4C2"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def slide_base():
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor(*_rgb(bg))
        return s

    def text(s, x, y, w, h, txt, size=14, bold=False, color=ink):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = str(txt)
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor(*_rgb(color))
        return tb

    def panel(s, x, y, w, h):
        from pptx.enum.shapes import MSO_SHAPE
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.adjustments[0] = 0.06
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor(*_rgb(card))
        shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    home = compute(db, user, dash, "home", filters).get("home") or {}
    s = slide_base()
    text(s, 0.7, 0.6, 9, 0.4, (spec.get("source") or {}).get("name", ""), 13, False, accent)
    text(s, 0.7, 1.0, 11.5, 1.6, spec.get("title") or dash.name, 44, True)
    text(s, 0.7, 2.3, 11.5, 0.6, spec.get("subtitle") or "", 18, False, muted)
    for i, k in enumerate((home.get("kpis") or [])[:4]):
        x = 0.7 + i * 3.05
        panel(s, x, 3.4, 2.85, 1.7)
        text(s, x + 0.25, 3.55, 2.4, 0.35, k["label"], 12, False, muted)
        text(s, x + 0.25, 3.9, 2.4, 0.7, _fmt(k["value"], k.get("unit") or "num") if k.get("unit") != "count" else f"{int(k['value'] or 0):,}", 30, True)
        text(s, x + 0.25, 4.6, 2.4, 0.35, k.get("note") or "", 11, False, muted)
    y = 5.5
    for d in (home.get("definitions") or [])[:3]:
        text(s, 0.7, y, 12, 0.3, d, 10, False, muted)
        y += 0.3
    for p in spec.get("pages") or []:
        if p.get("kind") == "home":
            continue
        out = compute(db, user, dash, p["id"], filters)
        blocks = [b for b in out.get("blocks") or [] if b.get("type") != "empty"]
        kpis = [b for b in blocks if b.get("type") == "kpi"]
        charts = [b for b in blocks if b.get("type") in ("bars", "line", "tiers")]
        tables = [b for b in blocks if b.get("type") in ("table", "list")]
        s = slide_base()
        text(s, 0.6, 0.35, 10, 0.6, f"{p['name']}", 28, True)
        text(s, 0.6, 0.95, 12, 0.4, p.get("blurb") or "", 13, False, muted)
        y0 = 1.5
        if kpis:
            w = 12.1 / max(len(kpis[:5]), 1)
            for i, k in enumerate(kpis[:5]):
                x = 0.6 + i * w
                panel(s, x, y0, w - 0.15, 1.25)
                text(s, x + 0.2, y0 + 0.1, w - 0.5, 0.3, k.get("title") or "", 11, False, muted)
                val = f"{int(k['value'] or 0):,}" if k.get("unit") == "count" else _fmt(k.get("value"), k.get("unit") or "num")
                text(s, x + 0.2, y0 + 0.4, w - 0.5, 0.5, val, 24, True, warn if k.get("tone") == "bad" else ink)
                text(s, x + 0.2, y0 + 0.88, w - 0.5, 0.3, " · ".join(c["text"] for c in (k.get("chips") or [])[:2]), 9, False, muted)
            y0 += 1.45
        for i, b in enumerate(charts[:2]):
            x = 0.6 + i * 6.1
            panel(s, x, y0, 5.95, 7.5 - y0 - 0.35)
            text(s, x + 0.2, y0 + 0.1, 5.5, 0.35, b.get("title") or "", 13, True)
            cd = CategoryChartData()
            items = b.get("rows") or b.get("points") or ([{"label": t["label"], "value": t["count"]} for t in b.get("tiers") or []])
            cd.categories = [_cap(str(it.get("label")))[:42] for it in items]
            unit = b.get("unit") or ("count" if b.get("type") == "tiers" else "num")
            cd.add_series(b.get("title") or "value", [it.get("value") or 0 for it in items])
            ctype = XL_CHART_TYPE.LINE_MARKERS if b.get("type") == "line" else XL_CHART_TYPE.BAR_CLUSTERED if b.get("type") == "bars" else XL_CHART_TYPE.COLUMN_CLUSTERED
            gf = s.shapes.add_chart(ctype, Inches(x + 0.15), Inches(y0 + 0.5), Inches(5.65), Inches(7.5 - y0 - 1.0), cd)
            ch = gf.chart
            ch.has_title = False
            ch.has_legend = False
            ch.font.size = Pt(10)
            ch.font.color.rgb = RGBColor(*_rgb(ink))
            nf = {"money": '"$"#,##0,"K"' if max((abs(it.get("value") or 0) for it in items), default=0) >= 10000 else '"$"#,##0',
                  "pct": '0.0"%"', "ratio": "0.00", "count": "#,##0"}.get(unit, "#,##0.0" if b.get("decimals") is None else "0.000")
            plot = ch.plots[0]
            plot.has_data_labels = len(items) <= 10
            if plot.has_data_labels:
                plot.data_labels.font.size = Pt(9)
                plot.data_labels.font.color.rgb = RGBColor(*_rgb(ink))
                plot.data_labels.number_format = nf
                plot.data_labels.number_format_is_linked = False
            ser = plot.series[0]
            if b.get("type") == "line":
                ser.format.line.color.rgb = RGBColor(*_rgb(accent))
                ser.format.line.width = Emu(28575)
                ser.smooth = False
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = RGBColor(*_rgb(warn if unit == "pct" and b.get("kind", "").startswith("o_") else accent))
                plot.gap_width = 60
                if b.get("type") == "bars":
                    try:
                        ch.category_axis.reverse_order = True
                    except Exception:  # noqa: BLE001
                        pass
            try:
                va = ch.value_axis
                if b.get("type") != "line" or unit in ("pct", "count", "money"):
                    va.minimum_scale = 0
                va.has_major_gridlines = True
                va.major_gridlines.format.line.color.rgb = RGBColor(*_rgb(muted))
                va.major_gridlines.format.line.width = Emu(3175)
                va.tick_labels.number_format = nf
                va.tick_labels.number_format_is_linked = False
                va.tick_labels.font.color.rgb = RGBColor(*_rgb(muted))
                va.format.line.fill.background()
                ch.category_axis.tick_labels.font.color.rgb = RGBColor(*_rgb(ink))
                ch.category_axis.format.line.color.rgb = RGBColor(*_rgb(muted))
            except Exception:  # noqa: BLE001
                pass
        if tables and len(charts) < 2:
            b = tables[0]
            x = 0.6 + (6.1 if charts else 0)
            w = 5.95 if charts else 12.1
            panel(s, x, y0, w, 7.5 - y0 - 0.35)
            text(s, x + 0.2, y0 + 0.1, w - 0.4, 0.35, b.get("title") or "", 13, True)
            if b.get("type") == "table":
                cols = (b.get("columns") or [])[:6]
                rows = (b.get("rows") or [])[:8] + ([b["total"]] if b.get("total") else [])
            else:
                cols = [{"key": "id", "label": b.get("id_label") or "Row"}, {"key": "score", "label": "Score", "unit": "ratio"}, {"key": "tags", "label": "Signals"}]
                rows = [{**r_, "tags": " · ".join(r_.get("tags") or [])} for r_ in (b.get("rows") or [])[:8]]
            tbl = s.shapes.add_table(len(rows) + 1, len(cols), Inches(x + 0.2), Inches(y0 + 0.55), Inches(w - 0.4), Inches(min(0.32 * (len(rows) + 1), 7.5 - y0 - 1))).table
            for j, c in enumerate(cols):
                cell = tbl.cell(0, j)
                cell.text = c["label"]
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(*_rgb(accent))
                cell.text_frame.paragraphs[0].font.size = Pt(10)
                cell.text_frame.paragraphs[0].font.bold = True
                cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255) if th["mode"] == "light" else RGBColor(*_rgb(bg))
            for i, r_ in enumerate(rows, start=1):
                for j, c in enumerate(cols):
                    v = r_.get(c["key"])
                    cell = tbl.cell(i, j)
                    cell.text = "" if v is None else (_fmt(v, c.get("unit")) if c.get("unit") and c.get("unit") != "count" and isinstance(v, (int, float)) else (f"{v:,}" if isinstance(v, int) else str(v)))
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(*_rgb(card))
                    cell.text_frame.paragraphs[0].font.size = Pt(9)
                    cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(*_rgb(ink))
        text(s, 0.6, 7.1, 12, 0.3, f"{(spec.get('source') or {}).get('name', '')} · {out.get('rows_in_view', 0):,} rows in view" +
             (" · " + ", ".join(out.get("applied") or []) if out.get("applied") else ""), 9, False, muted)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
