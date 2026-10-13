"""
GD360 Analyst (2026-10-13, R43 "Analyst brain").

Quick answer used to be one read-only SQL query per source, so any question
that needs a model, a formula, a rate nobody typed, a scenario, a
simulation, an optimisation, finance or statistics ended in "I can't answer
that from this data yet". The Analyst answers those the way a senior
analyst does:

  1. reads the data (or none at all - "simulate a coffee shop with random
     arrivals" needs no file), including the numbers the file states itself
     (a totals row: services/file_import.tidy_frame) and its formulas;
  2. the model writes a short plan and Python that computes the answer;
  3. GD360 runs it in the sandbox (services/sandbox.py) with a helper
     library of exact business maths (analyst_lib.py);
  4. checks it: the code must reproduce the file's own numbers before any
     scenario counts, and a run that errors or fails a check is sent back
     to the model with the error, up to three times;
  5. a writer turns ONLY the computed numbers into a decision-first answer;
     a number checker blocks any number that was not computed;
  6. a "scenario model" (inputs + safe formulas over the computed rows)
     lets the page recompute instantly when a rate changes - no model call.

Progress is written to run.plan["progress"] as it goes, so the page shows
the work live, the way a person watches an analyst think.
"""
from __future__ import annotations

import ast
import io
import itertools
import json
import math
import operator
import re
import statistics
import time
from datetime import datetime

import numpy as np
import pandas as pd
from sqlalchemy.orm import Session
from sqlalchemy.orm.attributes import flag_modified

from ... import models
from ...config import get_settings
from .. import ai_engine, data_access_rules, data_loader, synced_sources
from .analyst_lib import HELPER_DOC, HELPERS
from .analysis import table_payload
from .catalog import Catalog
from .numbers import allowed_values, extract_numbers, unsupported_numbers

settings = get_settings()

ANALYST_MARKER = "GD360 ANALYST"
WRITER_MARKER = "GD360 ANALYST WRITER"
MAX_ATTEMPTS = 3
MAX_TABLES = 6
MAX_ROWS = 300_000
LIVE_ROWS = 50_000

# Questions that need a model, a formula, a scenario or maths SQL can't hold.
_ROUTE = re.compile(
    r"\bwhat[\s-]*if\b|\bscenarios?\b|\bsimulat|\bmonte\s*carlo|\brandom (numbers?|data|sample)|\buse random\b|"
    r"\boptimi[sz]|\bmaximi[sz]e\b|\bminimi[sz]e\b|\bcheapest (option|number|plan|way|mix)|\bbest (number|mix|combination) of\b|"
    r"\bbreak[\s-]*even|\bsensitivity\b|\bnpv\b|\birr\b|\bpayback\b|\bpresent value\b|\bdiscount(ed)? cash|\bloan\b|\bemi\b|"
    r"\bprobability\b|\bexpected value\b|\bhypothesis\b|\bstatistically significant|\bp[\s-]?value\b|\bconfidence interval\b|"
    r"\bregression\b|\bcorrelation\b|\belasticity\b|\bqueue|\bbaristas?\b|\bstaffing\b|\bhow many (staff|people|servers|agents|cashiers|machines|trucks) (do|should|would)\b|"
    r"\bif (i|we) (add|hire|increase|decrease|raise|cut|double|halve|remove|reduce|open|close)\b|"
    r"\b(increase|decrease|raise|cut|reduce)\b[^.?!]*\bby\s+\d+\s*%[^.?!]*\b(what|how|would|impact|effect)\b|"
    r"\b(add|hire)\s+\d+\s+(more\s+)?\w+|"
    r"\b(labou?r|holding|ordering|waiting|opportunity|total daily|daily) cost\b|\bcost of (waiting|delay|holding)\b|"
    r"\beoq\b|\breorder point\b|\bsafety stock\b|\bnewsvendor\b|\blinear program|\bsample size\b|\ba/?b test\b",
    re.I,
)


def wants_analyst(question: str, catalog: Catalog, kind: str = "project") -> bool:
    """True when the question should go to the Analyst: no data at all, or a
    question that needs modelling. Guided Analysis keeps its step-by-step
    SQL plan."""
    if kind == "guided":
        return False
    if not catalog.sources:
        return True
    return bool(_ROUTE.search(question or ""))


# ---- progress ---------------------------------------------------------------

STAGES = [
    ("read", "Reading the data"),
    ("plan", "Working out the method"),
    ("run", "Running the analysis"),
    ("check", "Checking the numbers"),
    ("write", "Writing the answer"),
]


class _Progress:
    def __init__(self, db: Session, run: models.ProjectRun):
        self.db, self.run = db, run
        self.t0 = time.perf_counter()
        plan = dict(run.plan or {}) if isinstance(run.plan, dict) else {}
        plan["mode"] = "analyst"
        plan["progress"] = [{"id": k, "label": lbl, "status": "pending", "detail": None} for k, lbl in STAGES]
        run.plan = plan
        run.steps = []
        run.status = "running"
        run.started_at = run.started_at or datetime.utcnow()
        self._save()

    def set(self, stage: str, status: str, detail: str | None = None, **extra) -> None:
        plan = dict(self.run.plan or {})
        prog = [dict(p) for p in plan.get("progress") or []]
        for p in prog:
            if p["id"] == stage:
                p["status"] = status
                if detail is not None:
                    p["detail"] = detail
                p.update(extra)
                if status in ("done", "failed"):
                    p["at_s"] = round(time.perf_counter() - self.t0, 1)
        plan["progress"] = prog
        self.run.plan = plan
        flag_modified(self.run, "plan")
        self._save()

    def meta(self, **kv) -> None:
        plan = dict(self.run.plan or {})
        plan.update(kv)
        self.run.plan = plan
        flag_modified(self.run, "plan")
        self._save()

    def _save(self) -> None:
        try:
            self.db.commit()
        except Exception:  # noqa: BLE001
            self.db.rollback()


# ---- the data ---------------------------------------------------------------

def _py_name(name: str, taken: set[str]) -> str:
    base = re.sub(r"[^0-9a-zA-Z_]+", "_", (name or "").strip()).strip("_").lower() or "data"
    if base[0].isdigit():
        base = f"t_{base}"
    if base in ("pd", "np", "stats", "optimize", "math", "rng", "df", "tables", "result", "statistics", "itertools"):
        base = f"{base}_data"
    out, i = base, 2
    while out in taken:
        out, i = f"{base}_{i}", i + 1
    taken.add(out)
    return out


_STOP = {"the", "a", "an", "of", "and", "or", "to", "in", "on", "for", "by", "with", "what", "if", "i", "we", "is", "are",
         "my", "our", "how", "much", "many", "data", "file", "table", "sheet", "show", "give", "me", "this", "that", "it"}


def _words(text: str) -> set[str]:
    out = set()
    for w in re.findall(r"[a-z][a-z0-9]+", (text or "").lower()):
        if w not in _STOP and len(w) > 2:
            out.add(w[:-1] if w.endswith("s") and len(w) > 4 else w)
    return out


def relevant_sources(catalog: Catalog, question: str, limit: int = 3) -> list:
    """With one or two sources in scope they were chosen - use them. With
    many (every connected source, nothing picked), only the ones whose names,
    tables or columns the question touches; none when the question is
    self-contained ("simulate 100 customers with random arrivals")."""
    if len(catalog.sources) <= 2:
        return list(catalog.sources)
    q = _words(question)
    scored = []
    for pos, s in enumerate(catalog.sources):
        names = _words(s.name) | {w for t in s.tables for w in _words(t.name)}
        cols = {w for t in s.tables for c in t.columns for w in _words(str(c.get("name")))}
        score = 3 * len(q & names) + len(q & cols)
        if score >= 2:
            scored.append((-score, pos, s))
    scored.sort(key=lambda x: (x[0], x[1]))
    return [x[2] for x in scored[:limit]]


def load_frames(db: Session, user: models.User, catalog: Catalog, question: str = "") -> tuple[dict, list[dict]]:
    """{python name: DataFrame} for the tables in scope, plus a descriptor
    per table (source, rows, the file's own checks and tidy notes)."""
    frames: dict[str, pd.DataFrame] = {}
    info: list[dict] = []
    taken: set[str] = set()
    for src in relevant_sources(catalog, question):
        ds = db.get(models.DataSource, src.id)
        if ds is None:
            continue
        tables = [t for t in src.tables if t.rows != 0][: max(1, MAX_TABLES - len(frames))]
        for t in tables:
            if len(frames) >= MAX_TABLES:
                break
            note = None
            try:
                if src.mode == "synced":
                    df = synced_sources.load_table(db, ds.id, t.source_key)
                elif src.mode == "live":
                    df = data_loader.load_dataframe(ds, table=t.source_key or t.name, row_limit=LIVE_ROWS, db=db)
                    if len(df) >= LIVE_ROWS:
                        note = f"First {LIVE_ROWS:,} rows of a live table."
                else:
                    df = data_loader.load_dataframe(ds, table=t.source_key, db=db)
            except Exception as e:  # noqa: BLE001
                print(f"[analyst] could not load {src.name}/{t.name}: {e}")
                continue
            checks = list((df.attrs or {}).get("gd360_checks") or [])
            notes = list((df.attrs or {}).get("gd360_tidy") or [])
            try:
                df = data_access_rules.filter_dataframe_for_role(db, df, ds, user)
            except Exception as e:  # noqa: BLE001
                print(f"[analyst] access rules failed on {src.name}: {e}")
                continue
            if len(df) > MAX_ROWS:
                df = df.head(MAX_ROWS)
                note = f"First {MAX_ROWS:,} rows."
            label = src.name if len(src.tables) == 1 else f"{src.name} / {t.name}"
            name = _py_name(t.name if len(src.tables) > 1 else src.name, taken)
            frames[name] = df
            info.append({"name": name, "label": label, "source": src.name, "source_id": src.id, "kind": src.label,
                         "rows": int(len(df)), "columns": int(df.shape[1]), "checks": checks, "notes": notes,
                         "note": note, "formulas": _sheet_formulas(ds, t.source_key)})
    return frames, info


def _sheet_formulas(ds: models.DataSource, sheet: str | None) -> list[str]:
    """The first formulas of an Excel sheet ("D6 =SUM(D5+C6)", one per
    column) - they say what model the sheet is (a queue, a loan...)."""
    if ds.kind != "excel" or not ds.file_data:
        return []
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(ds.file_data), read_only=True, data_only=False)
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
        seen_cols: set[str] = set()
        out: list[str] = []
        headers: dict[str, str] = {}
        for row in ws.iter_rows(min_row=1, max_row=60):
            for c in row:
                v = c.value
                if v is None or not hasattr(c, "column_letter"):
                    continue
                col = c.column_letter
                if isinstance(v, str) and v.startswith("="):
                    if col not in seen_cols:
                        seen_cols.add(col)
                        out.append(f"{c.coordinate} {v}" + (f"   [{headers[col]}]" if col in headers else ""))
                elif isinstance(v, str) and col not in seen_cols:
                    headers[col] = (headers.get(col, "") + " " + v.strip()).strip()[:60]
            if len(out) >= 24:
                break
        wb.close()
        return out[:24]
    except Exception as e:  # noqa: BLE001
        print(f"[analyst] formulas skipped: {e}")
        return []


def _describe(name: str, df: pd.DataFrame, meta: dict) -> str:
    lines = [f'TABLE {name}  (from "{meta["label"]}", {meta["kind"]}) - {len(df):,} rows x {df.shape[1]} columns'
             + (f"  [{meta['note']}]" if meta.get("note") else "")]
    cols = []
    for c in list(df.columns)[:60]:
        s = df[c]
        if pd.api.types.is_numeric_dtype(s) and not pd.api.types.is_bool_dtype(s):
            nn = s.dropna()
            rng = f"{nn.min():g}..{nn.max():g}" if len(nn) else "empty"
            cols.append(f"{c!r} number {rng}")
        elif pd.api.types.is_datetime64_any_dtype(s):
            nn = s.dropna()
            cols.append(f"{c!r} date {str(nn.min())[:10]}..{str(nn.max())[:10]}" if len(nn) else f"{c!r} date")
        else:
            top = s.dropna().astype(str).value_counts().head(4).index.tolist()
            cols.append(f"{c!r} text, {s.nunique(dropna=True)} distinct e.g. {', '.join(t[:24] for t in top)}")
    lines.append("  columns: " + "; ".join(cols))
    with pd.option_context("display.max_columns", 40, "display.width", 220, "display.max_colwidth", 28):
        lines.append("  first rows:\n" + "\n".join("    " + x for x in df.head(8).to_string(index=False).splitlines()))
        if len(df) > 12:
            lines.append("  last rows:\n" + "\n".join("    " + x for x in df.tail(3).to_string(index=False, header=False).splitlines()))
    for ch in meta.get("checks") or []:
        if ch.get("column"):
            lines.append(f"  STATED BY THE FILE ITSELF (a totals row, already removed from the data): {ch['kind']} of {ch['column']!r} = {ch['value']:g}"
                         f"  (label: {ch.get('label')!r}) - reproduce this exactly before any scenario")
    for n in meta.get("notes") or []:
        lines.append(f"  reading note: {n}")
    if meta.get("formulas"):
        lines.append("  the sheet's formulas (first of each column): " + " | ".join(meta["formulas"]))
    return "\n".join(lines)


# ---- the model --------------------------------------------------------------

ANALYST_SYSTEM = f"""You are the {ANALYST_MARKER}: a brilliant senior analyst (operations research, statistics,
finance, economics and data science) inside GD360, a business analytics product used by
enterprise teams. You answer the person's question by writing Python that GD360 runs.

How to think, like the best analyst in the room:
- First understand what the data IS (a queue simulation, a ledger, a funnel, a time series...). The
  sheet's formulas tell you the model. Rebuild that model in code from the raw columns.
- If the file states its own numbers (a totals row), REPRODUCE them first and put that in result["checks"]
  before running any scenario. If you can't reproduce them you have misunderstood the data.
- What-if / scenarios: change only what was asked and keep everything else identical (same random
  numbers, same arrivals) so options are compared fairly. Include today's case as the baseline.
- Missing inputs (a wage, a cost per minute, a price, a discount rate, a distribution): NEVER refuse.
  Pick a sensible, commonly used default, say so in "assumptions", and expose it as a model input so the
  person can change it. Ask nothing - decide and show.
- No data, or the person asks for random numbers / a simulation: generate data in code with
  rng = np.random.default_rng(42) and clear, stated distributions; size it sensibly; show a sample table.
- Ambiguous wording ("add 2 baristas" = 2 in total, or 2 more?): choose the most likely reading, state it
  in "reading", and give the other in "other_readings" written as a full question the person could ask.
- Decide something. If the question is a choice, find the best option AND the range of inputs where it
  stays best (break-even). Show the options around the best one too (asked 2 and 3? also show 1 = today,
  and 4-5), so the person sees where the optimum is. If it is a forecast or estimate, give a range.
- Columns that ARE random draws (e.g. "Rand #" next to a derived value) are the simulation's random numbers:
  reuse them so your baseline matches the file; only draw new ones when asked for fresh randomness.
- Follow the conventions of the file (a class assignment, a finance model): same units, same rounding.
- Prefer exact methods; when you simulate, use a fixed seed and enough runs (1,000+) and report the spread.
- Be exact. Vectorize; the code must finish in under 20 seconds. Never print; build `result`.

Code rules: the tables are already loaded as pandas DataFrames, each as a variable named as shown and also
in the dict `tables`. NO import statements, no files, no network, no plotting, no classes. {HELPER_DOC}

Only when the question asks for FACTS about the person's business that the attached data does not contain and
that no reasonable assumption could stand in for (e.g. "sales by store location" when no column holds a location),
set result = {{"cannot": "one plain sentence on what is missing", "alternatives": ["up to 3 questions this data CAN
answer"]}}. A missing RATE, price or parameter is never a reason for "cannot" - assume it.

Otherwise your Python MUST end by setting `result` to a dict:
result = {{
  "key_numbers": {{"baseline_total_wait_min": 843, "best_option": 2, ...}},   # EVERY number the answer may quote,
                                  # including derived ones (savings, % changes, break-even values), as plain floats/ints
  "tables": [{{"title": "Every option, side by side", "df": options_df}}, ...],  # 1-4 tables, main table first,
                                  # readable column names ("Baristas", "Total wait (min)", "Daily cost ($)"), max 200 rows
  "checks": [{{"label": "1 barista reproduces the sheet's total wait", "expected": 843, "got": 843.0}}],
  "chart": {{"table": 0, "type": "bar|line|area|scatter|histogram", "x": "Baristas", "y": ["Daily cost ($)"],
            "title": "Daily cost by number of baristas"}},      # or None when a chart adds nothing
  "model": None or {{                                            # live what-if: optional, include it for scenario /
      "inputs": [{{"key": "wait_cost", "label": "Cost of a waiting minute", "value": 1.0, "min": 0, "max": 5,
                  "step": 0.05, "unit": "$"}}],                   # cost / rate questions so the person can move rates
      "rows": [{{"label": "1 barista", "values": {{"baristas": 1, "wait_min": 843}}}}, ...],
      "outputs": [{{"key": "waiting_cost", "label": "Waiting cost", "expr": "wait_min * wait_cost", "format": "money"}},
                  {{"key": "total", "label": "Daily cost", "expr": "waiting_cost + labor", "format": "money"}}],
      "pick": {{"output": "total", "goal": "min"}}                  # which row is best
  }}
}}
Model "expr" may use + - * / ** ( ), min(), max(), abs(), round(), numbers, row value names, input keys and
earlier output keys. formats: money | number | integer | percent | minutes | hours.

Reply in exactly this form - a JSON block, then one Python block:
```json
{{"title": "short title (max 60 chars)", "understanding": "one sentence: what the data is and what is asked",
  "reading": "how you read the question", "other_readings": ["full alternative question", ...],
  "method": "one sentence: the method", "assumptions": ["each assumption, one sentence"],
  "steps": ["3-6 short steps of your approach, plain English"]}}
```
```python
# the code
```"""


WRITER_SYSTEM = f"""You are the {WRITER_MARKER}. Write the answer a senior analyst gives an executive, from
the COMPUTED results only. Never invent or round to a new number: every number you write must appear in
the results (you may write it rounded the obvious way, e.g. 21.075 -> 21.1, 963.0 -> $963).

Lead with the decision or the direct answer in the headline (max 22 words, the key number in it). Then 2-4
short sentences that explain why, with the trade-off and the condition under which the answer would change
(break-even) when there is one. Say which assumptions matter. Plain words, no jargon, no markdown.
Then up to 3 follow-up questions that push the analysis further (sensitivity, risk, a bigger scenario).

Answer with JSON only:
{{"headline": "...", "answer": "...", "bullets": ["optional short point", "..."], "next_questions": ["...", "...", "..."]}}"""


def _history_text(history: list[dict], prev_code: str | None) -> str:
    lines = []
    if history:
        lines.append("EARLIER IN THIS THREAD:")
        for h in history[-4:]:
            lines.append(f"- Q: {h.get('question')}")
            if h.get("headline"):
                lines.append(f"  A: {h.get('headline')}")
    if prev_code:
        lines.append("The code that answered the last analyst question (reuse and extend it for a follow-up):\n```python\n"
                     + prev_code[:7000] + "\n```")
    return "\n".join(lines)


_JSON_BLOCK = re.compile(r"```json\s*(\{.*?\})\s*```", re.S)
_PY_BLOCK = re.compile(r"```(?:python|py)?\s*\n(.*?)```", re.S)


def parse_reply(text: str) -> tuple[dict, str]:
    meta: dict = {}
    m = _JSON_BLOCK.search(text or "")
    if m:
        try:
            meta = json.loads(m.group(1))
        except ValueError:
            meta = {}
    code = ""
    rest = _JSON_BLOCK.sub("", text or "")
    for block in _PY_BLOCK.findall(rest):
        if "result" in block and not block.strip().startswith("{"):
            code = block
    if not code:
        if "result" in rest:
            code = rest.replace("```", "")
    return (meta if isinstance(meta, dict) else {}), clean_code(code)


def clean_code(code: str) -> str:
    """Drops import lines (every allowed module is already in scope) and
    plotting calls, which have nothing to draw on."""
    out = []
    for line in (code or "").splitlines():
        s = line.strip()
        if re.match(r"^(import\s+\w|from\s+[\w.]+\s+import\s)", s):
            continue
        if re.match(r"^(plt|sns|fig|ax)\b", s) or s.startswith("print("):
            indent = line[: len(line) - len(line.lstrip())]
            out.append(f"{indent}pass")
            continue
        out.append(line)
    return "\n".join(out).strip()


def _call(messages: list[dict], max_tokens: int = 9000) -> str:
    model = (getattr(settings, "ANALYST_MODEL", "") or "").strip() or None
    return ai_engine._call_llm_resilient(messages, max_tokens=max_tokens, model_override=model)


# ---- running and checking ---------------------------------------------------

def _scope(frames: dict) -> dict:
    from scipy import optimize
    extra = dict(HELPERS)
    extra.update({"optimize": optimize, "math": math, "statistics": statistics, "itertools": itertools,
                  "rng": np.random.default_rng(42)})
    extra.update(frames)
    return extra


def _num(v):
    try:
        if isinstance(v, (bool, np.bool_)):
            return None
        f = float(v)
        return f if math.isfinite(f) else None
    except (TypeError, ValueError):
        return None


def normalize_result(raw) -> tuple[dict | None, list[str]]:
    """The code's result made safe to store and show, plus the problems
    that should send the code back for a fix."""
    problems: list[str] = []
    if not isinstance(raw, dict):
        return None, ["`result` was not set to a dict."]
    if raw.get("cannot"):
        return {"cannot": str(raw["cannot"])[:400],
                "alternatives": [str(x)[:140] for x in (raw.get("alternatives") or []) if str(x).strip()][:3]}, []
    keys = {}
    for k, v in (raw.get("key_numbers") or {}).items() if isinstance(raw.get("key_numbers"), dict) else []:
        f = _num(v)
        if f is not None:
            keys[str(k)[:60]] = f
    tables = []
    for t in (raw.get("tables") or [])[:4] if isinstance(raw.get("tables"), list) else []:
        if not isinstance(t, dict):
            continue
        df = t.get("df")
        if isinstance(df, pd.Series):
            df = df.reset_index()
        if isinstance(df, list):
            try:
                df = pd.DataFrame(df)
            except Exception:  # noqa: BLE001
                df = None
        if not isinstance(df, pd.DataFrame) or df.empty:
            continue
        df = df.copy()
        if not isinstance(df.index, pd.RangeIndex):
            df = df.reset_index()
        df.columns = [str(c) for c in df.columns]
        tables.append({"title": str(t.get("title") or "Result")[:90], **table_payload(df.head(200), limit=200), "rows_total": int(len(df))})
    if not tables and not keys:
        problems.append("result has no key_numbers and no tables - put the answer's numbers in result.")
    checks = []
    for c in (raw.get("checks") or [])[:8] if isinstance(raw.get("checks"), list) else []:
        if not isinstance(c, dict):
            continue
        exp_, got = _num(c.get("expected")), _num(c.get("got"))
        ok = exp_ is not None and got is not None and abs(exp_ - got) <= max(abs(exp_) * 0.005, 0.011)
        checks.append({"label": str(c.get("label") or "Check")[:120], "expected": exp_, "got": got, "ok": bool(ok)})
        if not ok:
            problems.append(f"Check failed: {c.get('label')} expected {exp_} but the code got {got}. Find the misunderstanding and fix it.")
    chart = raw.get("chart") if isinstance(raw.get("chart"), dict) else None
    model = None
    if isinstance(raw.get("model"), dict):
        model, mp = normalize_model(raw["model"])
        problems += mp
    return {"key_numbers": keys, "tables": tables, "checks": checks, "chart": chart, "model": model}, problems


# a tiny, safe expression evaluator for the scenario model (no eval)
_OPS = {ast.Add: operator.add, ast.Sub: operator.sub, ast.Mult: operator.mul, ast.Div: operator.truediv,
        ast.Pow: operator.pow, ast.USub: operator.neg, ast.UAdd: operator.pos}
_FUNCS = {"min": min, "max": max, "abs": abs, "round": round}


def safe_eval(expr: str, names: dict) -> float:
    def ev(n):
        if isinstance(n, ast.Expression):
            return ev(n.body)
        if isinstance(n, ast.Constant) and isinstance(n.value, (int, float)):
            return float(n.value)
        if isinstance(n, ast.Name):
            if n.id not in names:
                raise ValueError(f"unknown name {n.id}")
            return float(names[n.id])
        if isinstance(n, ast.BinOp) and type(n.op) in _OPS:
            a, b = ev(n.left), ev(n.right)
            if isinstance(n.op, ast.Div) and b == 0:
                return float("nan")
            return _OPS[type(n.op)](a, b)
        if isinstance(n, ast.UnaryOp) and type(n.op) in _OPS:
            return _OPS[type(n.op)](ev(n.operand))
        if isinstance(n, ast.Call) and isinstance(n.func, ast.Name) and n.func.id in _FUNCS and not n.keywords:
            return float(_FUNCS[n.func.id](*[ev(a) for a in n.args]))
        raise ValueError("unsupported expression")
    if len(expr) > 300:
        raise ValueError("expression too long")
    return ev(ast.parse(expr, mode="eval"))


_KEY = re.compile(r"^[A-Za-z_][A-Za-z0-9_]{0,40}$")


def normalize_model(m: dict) -> tuple[dict | None, list[str]]:
    problems: list[str] = []
    inputs = []
    for i in (m.get("inputs") or [])[:8]:
        if not isinstance(i, dict) or not _KEY.match(str(i.get("key") or "")):
            continue
        v = _num(i.get("value"))
        if v is None:
            continue
        lo = _num(i.get("min"))
        hi = _num(i.get("max"))
        lo = 0.0 if lo is None else min(lo, v)
        hi = (v * 4 if v > 0 else 10.0) if hi is None else max(hi, v)
        step = _num(i.get("step")) or (round((hi - lo) / 100, 4) or 0.01)
        inputs.append({"key": i["key"], "label": str(i.get("label") or i["key"])[:60], "value": v, "min": lo, "max": hi,
                       "step": step, "unit": str(i.get("unit") or "")[:8]})
    rows = []
    for r in (m.get("rows") or [])[:30]:
        if not isinstance(r, dict) or not isinstance(r.get("values"), dict):
            continue
        vals = {k: _num(v) for k, v in r["values"].items() if _KEY.match(str(k)) and _num(v) is not None}
        rows.append({"label": str(r.get("label") or f"Option {len(rows) + 1}")[:50], "values": vals})
    outputs = []
    for o in (m.get("outputs") or [])[:8]:
        if isinstance(o, dict) and _KEY.match(str(o.get("key") or "")) and isinstance(o.get("expr"), str):
            fmt = o.get("format") if o.get("format") in ("money", "number", "integer", "percent", "minutes", "hours") else "number"
            outputs.append({"key": o["key"], "label": str(o.get("label") or o["key"])[:50], "expr": o["expr"][:300], "format": fmt})
    if not (inputs and rows and outputs):
        return None, []
    base = {i["key"]: i["value"] for i in inputs}
    for r in rows:
        names = {**base, **r["values"]}
        comp = {}
        for o in outputs:
            try:
                comp[o["key"]] = safe_eval(o["expr"], {**names, **comp})
            except Exception as e:  # noqa: BLE001
                problems.append(f"model output {o['key']} = {o['expr']!r} failed for row {r['label']}: {e}. Use only the "
                                "row value names, input keys and earlier output keys.")
                return None, problems
        r["computed"] = comp
    pick = m.get("pick") if isinstance(m.get("pick"), dict) else {}
    pick = {"output": pick.get("output") if pick.get("output") in [o["key"] for o in outputs] else outputs[-1]["key"],
            "goal": "max" if pick.get("goal") == "max" else "min"}
    return {"inputs": inputs, "rows": rows, "outputs": outputs, "pick": pick}, problems


def _run_code(code: str, frames: dict) -> tuple[dict | None, str | None, list[str]]:
    from ..sandbox import run_sandboxed
    timeout = max(20, int(getattr(settings, "SANDBOX_TIMEOUT_SECONDS", 45) or 45))
    raw, err = run_sandboxed(code, dict(frames), timeout=timeout, extra=_scope(frames))
    if err:
        return None, err, []
    res, problems = normalize_result(raw)
    return res, None, problems


# ---- writing ----------------------------------------------------------------

def _results_text(res: dict, max_rows: int = 30) -> str:
    out = ["KEY NUMBERS: " + json.dumps(res.get("key_numbers") or {})[:4000]]
    for t in res.get("tables") or []:
        cols = [c["name"] for c in t["columns"]]
        out.append(f"TABLE {t['title']} ({t.get('rows_total', len(t['rows']))} rows): " + " | ".join(cols))
        for r in t["rows"][:max_rows]:
            out.append("  " + " | ".join(str(r.get(c)) for c in cols))
    for c in res.get("checks") or []:
        out.append(f"CHECK {c['label']}: expected {c['expected']}, got {c['got']} ({'ok' if c['ok'] else 'FAILED'})")
    m = res.get("model")
    if m:
        out.append("LIVE MODEL INPUTS: " + ", ".join(f"{i['label']} = {i['value']:g}{i['unit']}" for i in m["inputs"]))
        for r in m["rows"]:
            out.append(f"  {r['label']}: " + ", ".join(f"{k}={v:g}" for k, v in {**r['values'], **r['computed']}.items()))
    return "\n".join(out)


def _allowed(res: dict, meta: dict) -> list[float]:
    vals = list((res.get("key_numbers") or {}).values())
    m = res.get("model")
    if m:
        vals += [i["value"] for i in m["inputs"]]
        for r in m["rows"]:
            vals += list(r["values"].values()) + list(r["computed"].values())
    for c in res.get("checks") or []:
        vals += [x for x in (c["expected"], c["got"]) if x is not None]
    base = [v for v in vals if v is not None and math.isfinite(v)]
    derived = []
    small = base[:60]
    for a, b in itertools.combinations(small, 2):
        derived.append(a - b)
        if b:
            derived.append((a - b) / b * 100)
            derived.append(a / b * 100)
        if a:
            derived.append((b - a) / a * 100)
    derived += [v * 100 for v in base if abs(v) <= 1.5]
    text = " ".join(meta.get("assumptions") or []) + " " + str(meta.get("reading") or "")
    return base + derived + allowed_values([], res.get("tables") or [], text)


def write_answer(question: str, meta: dict, res: dict) -> dict:
    allowed = _allowed(res, meta)
    msgs = [{"role": "system", "content": WRITER_SYSTEM},
            {"role": "user", "content": f"QUESTION: {question}\nHOW IT WAS READ: {meta.get('reading') or ''}\n"
                                        f"METHOD: {meta.get('method') or ''}\nASSUMPTIONS: {json.dumps(meta.get('assumptions') or [])}\n\n"
                                        + _results_text(res)}]
    for attempt in range(2):
        try:
            raw = ai_engine._extract_json(_call(msgs, max_tokens=3000))
        except Exception as e:  # noqa: BLE001
            print(f"[analyst] writer failed: {e}")
            break
        draft = {"headline": str(raw.get("headline") or "").strip()[:240], "answer": str(raw.get("answer") or "").strip()[:1600],
                 "bullets": [str(b).strip()[:240] for b in (raw.get("bullets") or []) if str(b).strip()][:4],
                 "next_questions": [str(q).strip()[:160] for q in (raw.get("next_questions") or []) if str(q).strip()][:3]}
        text = " ".join([draft["headline"], draft["answer"], *draft["bullets"]])
        bad = unsupported_numbers(text, allowed)
        if draft["headline"] and not bad:
            return {**draft, "written_by": "model"}
        msgs += [{"role": "assistant", "content": json.dumps(raw)[:4000]},
                 {"role": "user", "content": "These numbers are not in the results: " + ", ".join(bad[:8])
                  + ". Rewrite using only numbers from the results. JSON only."}]
    return {**fallback_answer(question, meta, res), "written_by": "template"}


def _fmt(v: float) -> str:
    if abs(v) >= 100 or float(v).is_integer():
        return f"{v:,.0f}"
    return f"{v:,.2f}".rstrip("0").rstrip(".")


def fallback_answer(question: str, meta: dict, res: dict) -> dict:
    keys = res.get("key_numbers") or {}
    parts = [f"{k.replace('_', ' ')}: {_fmt(v)}" for k, v in list(keys.items())[:4]]
    m = res.get("model")
    headline = meta.get("title") or "Here is the analysis"
    if m:
        best = _best_row(m)
        if best:
            o = next(o for o in m["outputs"] if o["key"] == m["pick"]["output"])
            headline = f"{best['label']} is the best option: {o['label'].lower()} {_fmt(best['computed'][o['key']])}."
    return {"headline": headline, "answer": "; ".join(parts) + ("." if parts else ""), "bullets": [],
            "next_questions": [q for q in meta.get("other_readings") or []][:2]}


def _best_row(m: dict) -> dict | None:
    key = m["pick"]["output"]
    rows = [r for r in m["rows"] if math.isfinite(r["computed"].get(key, float("nan")))]
    if not rows:
        return None
    return (min if m["pick"]["goal"] == "min" else max)(rows, key=lambda r: r["computed"][key])


def _visual(res: dict) -> list[dict]:
    """The chart the code asked for, in the shape every GD360 chart reads -
    only when it was asked for and the columns exist."""
    ch = res.get("chart")
    if not ch or not res.get("tables"):
        return []
    try:
        t = res["tables"][int(ch.get("table") or 0)]
    except (IndexError, TypeError, ValueError):
        return []
    names = [c["name"] for c in t["columns"]]
    x = ch.get("x")
    ys = [y for y in (ch.get("y") or []) if y in names][:4] if isinstance(ch.get("y"), list) else ([ch.get("y")] if ch.get("y") in names else [])
    if x not in names or not ys:
        return []
    df = pd.DataFrame(t["rows"])[[x, *ys]]
    df[x] = df[x].astype(str)
    kind = {"bar": "bar", "line": "line", "area": "area", "scatter": "scatter", "histogram": "bar"}.get(str(ch.get("type")), "bar")
    payload = table_payload(df, limit=200)
    for c in payload["columns"]:
        c["role"] = "dimension" if c["name"] == x else "measure"
    return [{"type": "chart", "title": str(ch.get("title") or t["title"])[:90], "chart_type": kind, "tables": ["analyst"], **payload}]


# ---- the run ----------------------------------------------------------------

def _prev_code(db: Session, run: models.ProjectRun) -> str | None:
    prev = (db.query(models.ProjectRun)
            .filter(models.ProjectRun.conversation_id == run.conversation_id, models.ProjectRun.id != run.id,
                    models.ProjectRun.status == "done", models.ProjectRun.created_at <= run.created_at)
            .order_by(models.ProjectRun.created_at.desc()).first())
    a = ((prev.result or {}).get("analyst") or {}) if prev else {}
    return a.get("code")


def run(db: Session, run_row: models.ProjectRun, catalog: Catalog, user: models.User, history: list[dict],
        fallback_note: str | None = None) -> dict | None:
    """Answers the run. Returns {"cannot", "alternatives"} (and writes nothing) when the facts are simply not
    in the data - the caller then shows the closest questions the data CAN answer."""
    t0 = time.perf_counter()
    prog = _Progress(db, run_row)
    prog.set("read", "running")
    frames, info = load_frames(db, user, catalog, run_row.question)
    if frames:
        detail = "; ".join(f"{m['label']}: {m['rows']:,} rows" for m in info[:3])
        tidy = [n for m in info for n in m.get("notes") or []]
        if tidy:
            detail += " · " + tidy[0]
    else:
        detail = "No data attached - the analysis builds its own numbers"
    prog.set("read", "done", detail)

    question = run_row.question
    data_text = "\n\n".join(_describe(m["name"], frames[m["name"]], m) for m in info) or \
        "NO DATA IS ATTACHED. Answer from the question alone: generate data with rng where needed."
    user_msg = "\n\n".join(x for x in [
        f"TODAY: {datetime.utcnow().date().isoformat()}",
        "DATA\n" + data_text,
        _history_text(history, _prev_code(db, run_row)),
        (f"NOTE: a single SQL query could not answer this ({fallback_note[:300]}). Decide whether code with stated "
         "assumptions can - or whether the facts are simply not in the data (then use \"cannot\").") if fallback_note else "",
        f"QUESTION: {question}",
    ] if x)
    messages = [{"role": "system", "content": ANALYST_SYSTEM}, {"role": "user", "content": user_msg}]

    prog.set("plan", "running")
    meta, code, res, last_err = {}, "", None, None
    attempts: list[dict] = []
    for attempt in range(1, MAX_ATTEMPTS + 1):
        try:
            reply = _call(messages)
        except Exception as e:  # noqa: BLE001
            print(f"[analyst] model call failed: {e}")
            last_err = ai_engine.friendly_ai_error(e)
            break
        m2, code2 = parse_reply(reply)
        if m2:
            meta = {**meta, **m2}
        if not code2:
            last_err = "The model did not return code."
            messages += [{"role": "assistant", "content": reply[:6000]},
                         {"role": "user", "content": "Return the ```json block and ONE ```python block that sets `result`."}]
            continue
        code = code2
        if attempt == 1:
            prog.set("plan", "done", meta.get("method") or "Method chosen")
            prog.meta(title=(meta.get("title") or question)[:80], analyst_steps=[str(s)[:140] for s in (meta.get("steps") or [])][:6])
        prog.set("run", "running", f"Attempt {attempt}" if attempt > 1 else None)
        started = time.perf_counter()
        res, err, problems = _run_code(code, frames)
        attempts.append({"attempt": attempt, "ok": not err and not problems, "error": (err or "; ".join(problems))[:600] if (err or problems) else None,
                         "ms": int((time.perf_counter() - started) * 1000)})
        if err:
            last_err = err
            messages += [{"role": "assistant", "content": reply[:12000]},
                         {"role": "user", "content": f"Running your code failed:\n{err[-2500:]}\nFix it and reply in the same form (json block + python block)."}]
            continue
        prog.set("run", "done", f"Ran in {attempts[-1]['ms'] / 1000:.1f} s")
        prog.set("check", "running")
        if problems:
            last_err = "; ".join(problems)
            messages += [{"role": "assistant", "content": reply[:12000]},
                         {"role": "user", "content": "The result has problems:\n- " + "\n- ".join(problems[:6])
                          + "\nFix them and reply in the same form (json block + python block)."}]
            if attempt < MAX_ATTEMPTS:
                prog.set("check", "failed", problems[0][:200])
                continue
        break

    if res is None:
        _fail(db, run_row, prog, last_err or "The analysis could not be completed.", meta, code, attempts, t0)
        return None
    if res.get("cannot"):
        plan = dict(run_row.plan or {})
        plan.pop("mode", None)
        plan.pop("progress", None)
        run_row.plan = plan
        flag_modified(run_row, "plan")
        db.commit()
        return res
    checks = res.get("checks") or []
    ok_checks = [c for c in checks if c["ok"]]
    prog.set("check", "done" if all(c["ok"] for c in checks) else "failed",
             (f"{len(ok_checks)} of {len(checks)} checks match" + (f" - {ok_checks[0]['label']}" if ok_checks else "")) if checks
             else "Numbers computed directly - no stated totals to match")
    prog.set("write", "running")
    answer = write_answer(question, meta, res)
    prog.set("write", "done")

    data_used = [{"name": m["label"], "rows": m["rows"], "source_id": m["source_id"], "notes": m.get("notes") or [],
                  "checks": m.get("checks") or []} for m in info]
    run_row.result = {
        "answer": {"headline": answer["headline"], "answer": answer["answer"], "bullets": answer.get("bullets") or [],
                   "causes": [], "ruled_out": [], "next_questions": answer.get("next_questions") or [],
                   "written_by": answer.get("written_by"), "options": []},
        "analysis_type": "analyst",
        "analyst": {
            "title": (meta.get("title") or question)[:80], "understanding": meta.get("understanding"),
            "reading": meta.get("reading"), "other_readings": [str(x)[:160] for x in (meta.get("other_readings") or [])][:3],
            "method": meta.get("method"), "assumptions": [str(a)[:240] for a in (meta.get("assumptions") or [])][:8],
            "steps": [str(s)[:160] for s in (meta.get("steps") or [])][:6],
            "code": code, "attempts": attempts, "checks": checks, "key_numbers": res.get("key_numbers") or {},
            "tables": res.get("tables") or [], "model": res.get("model"), "data_used": data_used,
            "no_data": not frames, "fallback_note": fallback_note,
            "elapsed_s": round(time.perf_counter() - t0, 1),
        },
        "facts": [], "summary": {}, "kpis": [], "visuals": _visual(res),
        "warnings": [],
        "evidence": [{"id": f"a{i + 1}", "title": t["title"], "source": "Analysis", **{k: t[k] for k in ("columns", "rows", "truncated")}}
                     for i, t in enumerate(res.get("tables") or [])],
        "sources_used": sorted({m["source"] for m in info}),
        "queries": 0,
    }
    plan = dict(run_row.plan or {})
    plan["title"] = run_row.result["analyst"]["title"]
    plan["can_answer"] = True
    run_row.plan = plan
    flag_modified(run_row, "plan")
    run_row.status = "done"
    run_row.finished_at = datetime.utcnow()
    conv = db.get(models.Conversation, run_row.conversation_id)
    if conv and (not conv.title or conv.title in ("New analysis", "New project")):
        conv.title = run_row.result["analyst"]["title"]
    db.add(models.Message(conversation_id=run_row.conversation_id, role="assistant",
                          content=f"{answer['headline']}\n\n{answer['answer']}".strip()))
    db.commit()
    return None


def _fail(db: Session, run_row: models.ProjectRun, prog: _Progress, reason: str, meta: dict, code: str,
          attempts: list, t0: float) -> None:
    for st in ("plan", "run", "check", "write"):
        p = next((x for x in (run_row.plan or {}).get("progress") or [] if x["id"] == st), None)
        if p and p["status"] in ("pending", "running"):
            prog.set(st, "failed" if p["status"] == "running" else "pending")
    friendly = reason if len(reason) < 200 and "Traceback" not in reason else "The analysis code kept failing after three tries."
    run_row.status = "failed"
    run_row.error_message = f"I couldn't finish this analysis: {friendly} Try rephrasing, or add the missing numbers to the question."
    run_row.result = {"analyst": {"code": code, "attempts": attempts, "title": meta.get("title"), "method": meta.get("method"),
                                  "elapsed_s": round(time.perf_counter() - t0, 1)}}
    run_row.finished_at = datetime.utcnow()
    db.commit()


# ---- Excel with live formulas -----------------------------------------------

_XL_FUN = {"min": "MIN", "max": "MAX", "abs": "ABS", "round": "ROUND"}


def to_excel_formula(expr: str, refs: dict[str, str]) -> str:
    """A model formula in Excel: names become cell references, ** becomes ^."""
    out = []
    for tok in re.findall(r"\d+\.?\d*(?:e[-+]?\d+)?|\.\d+|[A-Za-z_][A-Za-z0-9_]*|\*\*|[-+*/(),]|\s+", expr):
        if tok.strip() == "":
            continue
        if tok == "**":
            out.append("^")
        elif re.match(r"^[A-Za-z_]", tok):
            if tok in refs:
                out.append(refs[tok])
            elif tok in _XL_FUN:
                out.append(_XL_FUN[tok])
            else:
                raise ValueError(tok)
        else:
            out.append(tok)
    return "=" + "".join(out)


def workbook(run_row: models.ProjectRun, inputs: dict | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    res = run_row.result or {}
    a = res.get("analyst") or {}
    ans = res.get("answer") or {}
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0F5132")
    inp_fill = PatternFill("solid", fgColor="FFF4D6")
    best_fill = PatternFill("solid", fgColor="D9F7E8")
    ws = wb.active
    ws.title = "Answer"
    ws["A1"] = ans.get("headline") or a.get("title") or "GD360 analysis"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = ans.get("answer") or ""
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
    ws.row_dimensions[3].height = 90
    r = 5
    if a.get("assumptions"):
        ws.cell(r, 1, "Assumptions").font = Font(bold=True)
        for x in a["assumptions"]:
            r += 1
            ws.cell(r, 1, f"- {x}")
        r += 2
    ws.cell(r, 1, f"Question: {run_row.question}").font = Font(italic=True, color="666666")

    m = a.get("model")
    if m:
        ws2 = wb.create_sheet("Scenarios")
        ws2["A1"] = "Change the yellow cells - every cost below recalculates."
        ws2["A1"].font = Font(italic=True, color="555555")
        refs: dict[str, str] = {}
        row = 3
        for i in m["inputs"]:
            v = (inputs or {}).get(i["key"], i["value"])
            ws2.cell(row, 1, i["label"])
            c = ws2.cell(row, 2, float(v))
            c.fill = inp_fill
            ws2.cell(row, 3, i.get("unit") or "")
            refs[i["key"]] = f"$B${row}"
            row += 1
        row += 1
        value_keys = list(m["rows"][0]["values"].keys()) if m["rows"] else []
        pretty = lambda k: re.sub(r"\bmin\b", "(min)", k.replace("_", " ")).capitalize()  # noqa: E731
        headers = ["Option"] + [pretty(k) for k in value_keys] + [o["label"] for o in m["outputs"]]
        for j, h in enumerate(headers, start=1):
            c = ws2.cell(row, j, h)
            c.font, c.fill = head, fill
        first = row + 1
        for k, rr in enumerate(m["rows"]):
            rown = first + k
            ws2.cell(rown, 1, rr["label"])
            local = dict(refs)
            for j, key in enumerate(value_keys, start=2):
                ws2.cell(rown, j, rr["values"].get(key))
                local[key] = f"{get_column_letter(j)}{rown}"
            for j, o in enumerate(m["outputs"], start=2 + len(value_keys)):
                col = get_column_letter(j)
                try:
                    ws2.cell(rown, j, to_excel_formula(o["expr"], local))
                except ValueError:
                    ws2.cell(rown, j, rr.get("computed", {}).get(o["key"]))
                local[o["key"]] = f"{col}{rown}"
                if o["format"] == "money":
                    ws2.cell(rown, j).number_format = "#,##0.00"
        pick_col = 2 + len(value_keys) + [o["key"] for o in m["outputs"]].index(m["pick"]["output"])
        last = first + len(m["rows"]) - 1
        lbl = get_column_letter(pick_col)
        fn = "MIN" if m["pick"]["goal"] == "min" else "MAX"
        ws2.cell(last + 2, 1, "Best option").font = Font(bold=True)
        ws2.cell(last + 2, 2, f"=INDEX(A{first}:A{last},MATCH({fn}({lbl}{first}:{lbl}{last}),{lbl}{first}:{lbl}{last},0))").fill = best_fill
        for j in range(1, len(headers) + 1):
            ws2.column_dimensions[get_column_letter(j)].width = 18 if j > 1 else 26

    for t in (a.get("tables") or [])[:4]:
        name = re.sub(r"[\[\]\*\?/\\:]", " ", t.get("title") or "Table")[:31] or "Table"
        if name in wb.sheetnames:
            name = (name[:28] + f" {len(wb.sheetnames)}")[:31]
        wst = wb.create_sheet(name)
        cols = [c["name"] for c in t["columns"]]
        for j, h in enumerate(cols, start=1):
            c = wst.cell(1, j, h)
            c.font, c.fill = head, fill
            wst.column_dimensions[get_column_letter(j)].width = max(12, min(40, len(h) + 4))
        for i, rr in enumerate(t["rows"], start=2):
            for j, h in enumerate(cols, start=1):
                wst.cell(i, j, rr.get(h))

    wm = wb.create_sheet("Method")
    lines = [("Method", a.get("method") or "")] + [(f"Step {i + 1}", s) for i, s in enumerate(a.get("steps") or [])]
    lines += [(f"Check: {c['label']}", f"expected {c['expected']}, got {c['got']} ({'ok' if c['ok'] else 'did not match'})") for c in a.get("checks") or []]
    for i, (k, v) in enumerate(lines, start=1):
        wm.cell(i, 1, k).font = Font(bold=True)
        wm.cell(i, 2, v)
    n = len(lines) + 2
    wm.cell(n, 1, "Code").font = Font(bold=True)
    for i, line in enumerate((a.get("code") or "").splitlines(), start=n + 1):
        wm.cell(i, 1, line).font = Font(name="Consolas", size=9)
    wm.column_dimensions["A"].width = 40
    wm.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
