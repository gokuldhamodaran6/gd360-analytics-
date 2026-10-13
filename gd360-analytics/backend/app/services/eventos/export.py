"""The A to Z workbook: every section of the event as its own sheet, with
the truth label of each number. Personal data is masked unless asked."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ... import models
from . import views
from .ctx import METRICS, Ctx

SECTIONS = {
    "A": "Summary and success score", "B": "Expectations vs actual", "C": "Plan and tasks", "D": "Budget and P&L",
    "E": "Tracking and channels", "F": "Registrations and demand", "G": "Posts before", "H": "Event day",
    "I": "Sessions and livestream", "J": "Sponsors", "K": "Posts after", "L": "Follow-up", "M": "Pipeline and impact",
    "N": "ROI and confidence", "O": "Learnings", "P": "Raw data (people, leads, posts)",
}
HEAD = PatternFill("solid", fgColor="0F5C46")


def _mask_email(e: str | None) -> str | None:
    if not e or "@" not in e:
        return e
    u, d = e.split("@", 1)
    return (u[:1] + "•••@" + d) if u else e


def _mask_name(n: str | None) -> str | None:
    if not n:
        return n
    parts = n.split()
    return parts[0] + (" " + parts[-1][:1] + "." if len(parts) > 1 else "")


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], widths: list[int] | None = None):
    ws = wb.create_sheet(title[:31])
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([("" if x is None else x) for x in r])
    for n, wdt in enumerate(widths or [max(12, min(48, len(h) + 4)) for h in header], start=1):
        ws.column_dimensions[get_column_letter(n)].width = wdt
    ws.freeze_panes = "A2"
    return ws


def workbook(ctx: Ctx, sections: list[str] | None = None, mask: bool = True) -> bytes:
    want = set(sections or SECTIONS)
    rep = views.report(ctx)
    v = rep["values"]
    sc = rep["score"]
    wb = Workbook()
    ws = wb.active
    ws.title = "A Summary"
    head = rep["event"]
    rows = [["Event", ctx.i.title], ["Format", head["format_label"]], ["Starts", head["windows"]["start_local"]],
            ["Ends", head["windows"]["end_local"]], ["Where", ctx.i.location or ""], ["Phase now", head["phase_now"]],
            ["Event Success Score", sc["now"]["score"]], ["Confidence (share backed by counted/matched numbers)", f"{sc['now']['confidence']}%" if sc["now"]["confidence"] is not None else ""],
            ["Predicted score (median)", (sc["predicted"] or {}).get("median")], ["Goals signed", f"{sc['signed']} of {sc['total']}"],
            [], ["Number", "Value", "Truth"]]
    for k, (label, unit, truth, *_r) in METRICS.items():
        if v.get(k) is not None:
            rows.append([label, v.get(k), truth])
    for r in rows:
        ws.append(r)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 30
    ws["A1"].font = Font(bold=True, size=14)
    if "B" in want:
        _sheet(wb, "B Expectations", ["Owner", "Goal", "Measured as", "Target", "Low end", "Good end", "Weight", "Actual", "Truth",
                                      "Score", "Signed", "Signed by"],
               [[g["owner_name"], g["label"], g["metric_label"], g["target"], g["low"], g["high"], g["weight"], g["actual"], g["truth"],
                 g["score"], "yes" if g["signed"] else "no", g["signed_by"]] for g in sc["goals"]],
               [20, 36, 26, 12, 12, 12, 8, 14, 10, 8, 8, 18])
    if "C" in want:
        tasks = ctx.db.query(models.InitiativeTask).filter(models.InitiativeTask.initiative_id == ctx.i.id).order_by(models.InitiativeTask.due_on).all()
        _sheet(wb, "C Plan", ["Task", "Owner", "Due", "Status", "Closed by"],
               [[t.title, t.owner_name, t.due_on, t.status, (t.approval or {}).get("state") or ""] for t in tasks], [60, 18, 12, 10, 14])
    if "D" in want or "N" in want:
        rows = [["Cost", c["label"], c["amount"]] for c in rep["costs"]] + [["Income", x["label"], x["amount"]] for x in rep["income"]]
        rows += [[], ["Total cost", "", v.get("total_cost")], ["Income", "", v.get("income")], ["Cost per attendee", "", v.get("cost_per_attendee")],
                 ["Pipeline sourced", "", v.get("pipeline_sourced")], ["Expected revenue (win rate × sourced)", "", v.get("expected_revenue")],
                 ["ROI % (modelled)", "", v.get("roi_pct")]]
        _sheet(wb, "D Budget and PnL", ["Type", "Line", f"Amount ({rep['currency']})"], rows, [14, 48, 18])
    if "E" in want:
        _sheet(wb, "E Channels", ["Channel", "Registrations", "ICP %", "Came", "Show-up %", "Spend", "Cost per ICP registration"],
               [[c["channel"], c["registrations"], c["icp_share"], c["came"], c["show_rate"], round(c["spend"], 2), c["cost_per_icp"]]
                for c in rep["channels"]], [26, 14, 10, 10, 12, 14, 24])
        ads = ctx.ads()
        _sheet(wb, "E Ad campaigns", ["Platform", "Campaign", "Matched by", "Spend", "Clicks", "Impressions"],
               [[c["platform"], c["campaign"], c["how"], round(c["spend"], 2), c["clicks"], c["impressions"]] for c in ads["campaigns"]],
               [14, 48, 14, 14, 12, 14])
    if "F" in want:
        _sheet(wb, "F Daily", ["Date", "Phase", "Registrations", "Check-ins", "Posts", "Ad spend"],
               [[r["date"], r["phase"], r["registrations"], r["checkins"], r["mentions"], r["spend"]] for r in ctx.daily()], [12, 10, 14, 12, 10, 12])
    if "G" in want or "K" in want:
        for ph, tag in (("before", "G"), ("after", "K")):
            if tag not in want:
                continue
            ms = ctx.mentions_in(ph)
            _sheet(wb, f"{tag} Posts {ph}", ["Posted", "Platform", "Author", "Who", "Matched by", "Confidence", "Reach", "Engagements",
                                              "Sentiment", "Link", "Text"],
                   [[m.posted_at.isoformat() if m.posted_at else "", m.platform, m.author, m.author_kind, m.matched_by, m.confidence,
                     m.reach, m.engagements, m.sentiment, m.url, (m.text or "")[:500]] for m in ms],
                   [18, 10, 22, 10, 12, 10, 10, 12, 10, 36, 60])
    if "H" in want:
        _sheet(wb, "H Check-ins", ["Checked in", "Name", "Company", "Tier", "Walk-in", "Gate", "By", "ICP"],
               [[t.checked_in_at.isoformat() if t.checked_in_at else "", _mask_name(t.name) if mask else t.name, t.company, t.tier,
                 "yes" if t.walk_in else "", t.gate, t.checked_in_by, "yes" if t.icp else ("no" if t.icp is False else "")]
                for t in sorted(ctx.came, key=lambda t: t.checked_in_at or ctx.now)], [20, 22, 26, 10, 8, 10, 16, 6])
    if "I" in want:
        _sheet(wb, "I Sessions", ["Session", "Speaker", "Room", "Starts", "Capacity", "In room", "Fill %", "Online peak", "Rating"],
               [[s["title"], s["speaker"], s["room"], s["starts_at"], s["capacity"], s["in_room"], s["fill"], s["online_peak"], s["rating"]]
                for s in views.sessions_rows(ctx)], [40, 22, 14, 18, 10, 10, 8, 12, 8])
    if "J" in want:
        sp = views.sponsors(ctx)
        _sheet(wb, "J Sponsors", ["Sponsor", "Tier", "Booth", "Leads", "Goal", "% of goal", "Hot leads", "Booth visits", "Avg dwell (min)", "Renewal"],
               [[s["name"], s["tier"], s["booth"], s["leads"], s["goal"], s["pct"], s["hot"], s["visits"], s["dwell_min"], s["renewal"]]
                for s in sp["sponsors"]], [24, 10, 8, 8, 8, 10, 10, 12, 14, 10])
    if "L" in want:
        acts = ctx.db.query(models.EventAction).filter(models.EventAction.initiative_id == ctx.i.id).all()
        _sheet(wb, "L Follow-up and agent", ["Agent", "Suggestion", "State", "Result", "Decided by"],
               [[a.agent, a.title, a.state, a.result, a.decided_by] for a in acts], [14, 60, 12, 50, 16])
    if "M" in want:
        d = rep["deals"]
        _sheet(wb, "M Pipeline", ["Deal", "Company matched", "Kind", "Stage", "Amount", "Won", "Created", "How"],
               [[x["name"], x.get("company"), x["kind"], x.get("stage"), x["amount"], "yes" if x.get("won") else "", x.get("created_at"), x.get("how")]
                for x in d["deals"]], [40, 20, 12, 16, 14, 6, 20, 22])
    if "O" in want:
        _sheet(wb, "O Learnings", ["Kind", "Learning", "Proof"], [[l.get("kind"), l.get("title"), l.get("proof")] for l in rep["learnings"]], [10, 60, 60])
    if "P" in want:
        _sheet(wb, "P People", ["Ticket", "Name", "Email", "Company", "Title", "Tier", "Source", "Status", "Registered", "Checked in",
                                "Walk-in", "ICP", "Show-up probability", "Consent"],
               [[t.code, _mask_name(t.name) if mask else t.name, _mask_email(t.email) if mask else t.email, t.company, t.title, t.tier,
                 t.source, t.status, t.registered_at.isoformat() if t.registered_at else "", t.checked_in_at.isoformat() if t.checked_in_at else "",
                 "yes" if t.walk_in else "", "yes" if t.icp else ("no" if t.icp is False else ""), t.show_prob, "yes" if t.consent else ""]
                for t in ctx.tickets], [10, 22, 26, 24, 24, 8, 16, 12, 20, 20, 8, 6, 12, 8])
        _sheet(wb, "P Leads", ["Sponsor", "Name", "Email", "Company", "Title", "Rating", "Interests", "Note", "Scanned by", "At"],
               [[next((s.name for s in ctx.sponsors if s.id == l.sponsor_id), ""), _mask_name(l.name) if mask else l.name,
                 _mask_email(l.email) if mask else l.email, l.company, l.title, l.rating, ", ".join(l.interests or []), l.note, l.scanned_by,
                 l.created_at.isoformat() if l.created_at else ""] for l in ctx.leads], [18, 20, 24, 22, 22, 8, 26, 36, 14, 20])
        sv = ctx.by_kind.get("survey", [])
        _sheet(wb, "P Survey", ["Answered", "Score (0-10)", "Comment"], [[x.occurred_at.isoformat() if x.occurred_at else "", x.value, x.text] for x in sv],
               [20, 12, 80])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
